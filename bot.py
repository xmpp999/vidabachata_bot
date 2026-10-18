"""
Telegram-бот для продажи билетов на вечеринки Vida Bachata.
Хранилище: Supabase (PostgreSQL).
"""

import asyncio
import csv
import http.server
import io
import json
import os
import re
import socketserver
import threading
import urllib.request
import uuid
import zipfile
from datetime import datetime, timedelta
from typing import Optional

import dateparser
import qrcode
from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart, Command, CommandObject
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    Message, CallbackQuery, BufferedInputFile,
    InlineKeyboardMarkup, InlineKeyboardButton,
    ReplyKeyboardMarkup, KeyboardButton,
    ReplyKeyboardRemove,
)
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from supabase import create_client, Client
from yookassa import Configuration, Payment
from nalogo import Client as NalogClient

# ================== КОНФИГ ==================
load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN", "")
ADMIN_ID = int(os.getenv("ADMIN_ID", "0"))
BOT_USERNAME = os.getenv("BOT_USERNAME", "your_bot")

SUPABASE_URL = os.getenv("SUPABASE_URL", "")
SUPABASE_SERVICE_KEY = os.getenv("SUPABASE_SERVICE_KEY", "")

# ЮKassa
YOOKASSA_SHOP_ID = os.getenv("YOOKASSA_SHOP_ID", "")
YOOKASSA_SECRET_KEY = os.getenv("YOOKASSA_SECRET_KEY", "")

Configuration.account_id = YOOKASSA_SHOP_ID
Configuration.secret_key = YOOKASSA_SECRET_KEY

# ================== БОТ ==================
bot = Bot(BOT_TOKEN)
dp = Dispatcher()

# ================== SUPABASE ==================
supabase: Client = create_client(SUPABASE_URL, SUPABASE_SERVICE_KEY)

# ================== ХРАНИЛИЩЕ ==================
users = {}
username_index = {}
pending_invites = {}
tickets = {}
orders = []
events = []
blacklist = {}
reg_attempts = []
pending_payments = {}
processed_payments = set()
checkers = {}
ticket_logs = []
ticket_transfers = []

settings = {
    "support": os.getenv("SUPPORT", "@your_support"),
    "yoomoney_link": os.getenv("YOOMONEY_LINK", "https://yoomoney.ru/to/0000000000000"),
    "sales_stop_hours": 1,
    "auto_burn_hours": 6,
    "archive_after_hours": 6,
    "ticket_archive_days": 30,
}




# ================== ТОЧЕЧНОЕ СОХРАНЕНИЕ ==================

def save_user(uid: int):
    """Сохраняет одного пользователя."""
    u = users.get(uid)
    if not u:
        return
    try:
        supabase.table("users").upsert({
            "user_id": uid,
            "username": u.get("username", ""),
            "email": u.get("email", ""),
            "name": u.get("name", ""),
            "surname": u.get("surname", ""),
            "gender": u.get("gender", ""),
            "phone": u.get("phone", ""),
            "registered_at": u.get("registered_at", ""),
        }).execute()
    except Exception as e:
        print(f"[SUPABASE] save_user({uid}) ошибка: {e}")


def save_event(event_id: str):
    """Сохраняет одно событие."""
    e = next((x for x in events if str(x["id"]) == str(event_id)), None)
    if not e:
        return
    try:
        supabase.table("events").upsert({
            "id": e.get("id", ""),
            "title": e.get("title", ""),
            "date_text": e.get("date_text", ""),
            "start_dt": e.get("start_dt", ""),
            "place": e.get("place", ""),
            "description": e.get("description", ""),
            "price_single": e.get("price_single", 0),
            "price_pair": e.get("price_pair", 0),
            "photo_id": e.get("photo_id"),
            "created_at": e.get("created_at", ""),
            "active": e.get("active", True),
            "archived": e.get("archived", False),
            "archived_at": e.get("archived_at"),
        }).execute()
    except Exception as ex:
        print(f"[SUPABASE] save_event({event_id}) ошибка: {ex}")


def save_ticket(code: str):
    """Сохраняет один билет."""
    t = tickets.get(code)
    if not t:
        return
    try:
        supabase.table("tickets").upsert({
            "code": code,
            "user_id": t.get("user_id"),
            "type": t.get("type", ""),
            "holder": t.get("holder", ""),
            "used": t.get("used", False),
            "used_at": t.get("used_at", ""),
            "used_by": t.get("used_by", ""),
            "auto_used": t.get("auto_used", False),
            "archived_at": t.get("archived_at"),
            "event_id": t.get("event_id", ""),
            "event_title": t.get("event_title", ""),
            "event_date": t.get("event_date", ""),
            "issued_at": t.get("issued_at", ""),
            "partner_code": t.get("partner_code", ""),
        }).execute()
        
    except Exception as e:
        print(f"[SUPABASE] save_ticket({code}) ошибка: {e}")


def save_order(order: dict):
    """Сохраняет один заказ."""
    try:
        supabase.table("orders").upsert(order).execute()
    except Exception as e:
        print(f"[SUPABASE] save_order ошибка: {e}")

def save_ticket_transfer(transfer: dict):
    """Сохраняет запись о передаче билета."""
    try:
        supabase.table("ticket_transfers").insert(transfer).execute()
    except Exception as e:
        print(f"[SUPABASE] save_ticket_transfer ошибка: {e}")


def save_checker(chk_id: str):
    """Сохраняет одного checker'а (точечно, без delete+insert)."""
    c = checkers.get(chk_id)
    if not c:
        return
    try:
        row = {
            "user_id": c.get("user_id"),
            "username": c.get("username", ""),
            "phone": c.get("phone", ""),
            "name": c.get("name", ""),
            "event_id": c.get("event_id"),
            "added_by": c.get("added_by"),
            "added_at": c.get("added_at", ""),
            "active": c.get("active", True),
        }
        if c.get("db_id"):
            supabase.table("checkers").update(row).eq("id", c["db_id"]).execute()
        else:
            resp = supabase.table("checkers").insert(row).execute()
            if resp.data:
                c["db_id"] = resp.data[0]["id"]
    except Exception as e:
        print(f"[SUPABASE] save_checker({chk_id}) ошибка: {e}")


def delete_checker(chk_id: str):
    """Удаляет checker'а из БД (вызывать ДО удаления из памяти)."""
    c = checkers.get(chk_id)
    if not c:
        return
    db_id = c.get("db_id")
    if not db_id:
        return
    try:
        supabase.table("checkers").delete().eq("id", db_id).execute()
    except Exception as e:
        print(f"[SUPABASE] delete_checker({chk_id}) ошибка: {e}")




# ========== КОНЕЦ ЧАСТИ 1 ==========


def save_data():
    """
    Аварийное полное сохранение (только при старте/миграции).
    В обычной работе используйте точечные save_* функции.
    """
    try:
        if users:
            supabase.table("users").upsert([
                {
                    "user_id": uid,
                    "username": u.get("username", ""),
                    "email": u.get("email", ""),
                    "name": u.get("name", ""),
                    "surname": u.get("surname", ""),
                    "gender": u.get("gender", ""),
                    "phone": u.get("phone", ""),
                    "registered_at": u.get("registered_at", ""),
                }
                for uid, u in users.items()
            ]).execute()

        if events:
            supabase.table("events").upsert([
                {
                    "id": e.get("id", ""),
                    "title": e.get("title", ""),
                    "date_text": e.get("date_text", ""),
                    "start_dt": e.get("start_dt", ""),
                    "place": e.get("place", ""),
                    "description": e.get("description", ""),
                    "price_single": e.get("price_single", 0),
                    "price_pair": e.get("price_pair", 0),
                    "photo_id": e.get("photo_id"),
                    "created_at": e.get("created_at", ""),
                    "active": e.get("active", True),
                    "archived": e.get("archived", False),
                    "archived_at": e.get("archived_at"),
                }
                for e in events
            ]).execute()

        if tickets:
            supabase.table("tickets").upsert([
                {
                    "code": code,
                    "user_id": t.get("user_id"),
                    "type": t.get("type", ""),
                    "holder": t.get("holder", ""),
                    "used": t.get("used", False),
                    "used_at": t.get("used_at", ""),
                    "used_by": t.get("used_by", ""),
                    "auto_used": t.get("auto_used", False),
                    "archived_at": t.get("archived_at"),
                    "event_id": t.get("event_id", ""),
                    "event_title": t.get("event_title", ""),
                    "event_date": t.get("event_date", ""),
                    "issued_at": t.get("issued_at", ""),
                }
                for code, t in tickets.items()
            ]).execute()

        if orders:
            supabase.table("orders").upsert(orders).execute()

        if ticket_logs:
            for log in ticket_logs:
                try:
                    supabase.table("ticket_logs").insert(log).execute()
                except Exception:
                    pass
            ticket_logs.clear()

    except Exception as e:
        print(f"[SUPABASE] Ошибка полного сохранения: {e}")


def load_data():
    """Загружает всё из Supabase в память."""
    global events, orders, tickets, users, username_index, blacklist, reg_attempts, checkers, ticket_logs
    ticket_logs = []
    pending_payments.clear()

    try:
        resp = supabase.table("users").select("*").execute()
        users = {}
        username_index = {}
        for row in resp.data:
            uid = int(row["user_id"])
            users[uid] = {
                "username": row.get("username") or "",
                "email": row.get("email") or "",
                "name": row.get("name") or "",
                "surname": row.get("surname") or "",
                "gender": row.get("gender") or "",
                "phone": row.get("phone") or "",
                "registered_at": row.get("registered_at") or "",
            }
            uname = (row.get("username") or "").lstrip("@").strip().lower()
            if uname:
                username_index[uname] = uid

        resp = supabase.table("events").select("*").execute()
        events = []
        for row in resp.data:
            events.append({
                "id": row.get("id") or "",
                "title": row.get("title") or "",
                "date_text": row.get("date_text") or "",
                "start_dt": row.get("start_dt") or "",
                "place": row.get("place") or "",
                "description": row.get("description") or "",
                "price_single": row.get("price_single", 0),
                "price_pair": row.get("price_pair", 0),
                "photo_id": row.get("photo_id"),
                "created_at": row.get("created_at") or "",
                "active": row.get("active", True),
                "archived": row.get("archived", False),
                "archived_at": row.get("archived_at"),
            })

        resp = supabase.table("tickets").select("*").execute()
        tickets = {}
        for row in resp.data:
            tickets[row["code"]] = {
                "user_id": row.get("user_id"),
                "type": row.get("type") or "",
                "holder": row.get("holder") or "",
                "used": row.get("used", False),
                "used_at": row.get("used_at") or "",
                "used_by": row.get("used_by") or "",
                "auto_used": row.get("auto_used", False),
                "archived_at": row.get("archived_at"),
                "event_id": row.get("event_id") or "",
                "event_title": row.get("event_title") or "",
                "event_date": row.get("event_date") or "",
                "issued_at": row.get("issued_at") or "",
                "partner_code": row.get("partner_code") or "",
            }

        resp = supabase.table("orders").select("*").execute()
        orders = resp.data or []

        resp = supabase.table("blacklist").select("*").execute()
        blacklist = {}
        for row in resp.data:
            v = (row.get("value") or "").lower().lstrip("@")
            v = v.replace(" ", "").replace("-", "").replace("+", "")
            key = f"{row.get('type')}:{v}"
            blacklist[key] = {
                "type": row.get("type") or "",
                "value": row.get("value") or "",
                "reason": row.get("reason") or "",
                "added_at": row.get("added_at") or "",
            }

        resp = supabase.table("reg_attempts").select("*").order("id", desc=False).limit(500).execute()
        reg_attempts = resp.data or []

        resp = supabase.table("checkers").select("*").execute()
        checkers = {}
        for row in resp.data:
            chk_id = f"chk_{row['id']}"
            checkers[chk_id] = {
                "id": chk_id,
                "db_id": row["id"],
                "user_id": row.get("user_id"),
                "username": (row.get("username") or "").lstrip("@").lower(),
                "phone": row.get("phone") or "",
                "name": row.get("name") or "",
                "event_id": row.get("event_id"),
                "added_by": row.get("added_by"),
                "added_at": row.get("added_at") or "",
                "active": row.get("active", True),
            }

        print(
            f"✅ Загружено из Supabase: {len(events)} событий, "
            f"{len(orders)} заказов, {len(users)} юзеров, "
            f"{len(tickets)} билетов, {len(blacklist)} в ЧС, "
            f"{len(reg_attempts)} попыток"
        )
        print(f"   └ Checkers: {len(checkers)}")
    except Exception as e:
        print(f"[SUPABASE] Ошибка загрузки: {e}")

def load_transfers():
    """Загружает историю передач из Supabase."""
    global ticket_transfers
    try:
        resp = supabase.table("ticket_transfers").select("*").order("transferred_at", desc=True).limit(500).execute()
        ticket_transfers = resp.data or []
    except Exception as e:
        print(f"[SUPABASE] load_transfers ошибка: {e}")
        ticket_transfers = []


# ================== ЮKASSA ==================
def create_yookassa_payment(amount: int, description: str, user_id: int, metadata: dict) -> dict:
    """Создаёт платёж в ЮKassa."""
    idempotence_key = str(uuid.uuid4())

    payment_data = {
        "amount": {
            "value": f"{amount}.00",
            "currency": "RUB"
        },
        "capture": True,
        "confirmation": {
            "type": "redirect",
            "return_url": "https://t.me/" + BOT_USERNAME
        },
        "description": description,
        "metadata": {
            "user_id": user_id,
            **metadata
        }
    }

    try:
        payment = Payment.create(payment_data, idempotence_key)
        return {
            "payment_id": payment.id,
            "confirmation_url": payment.confirmation.confirmation_url,
            "status": payment.status
        }
    except Exception as e:
        print(f"[YOOKASSA] Ошибка создания платежа: {e}")
        return {}


# ========== КОНЕЦ ЧАСТИ 2 ==========


def verify_payment_with_yookassa(payment_id: str) -> bool:
    """Проверяет платёж через API ЮKassa."""
    try:
        payment = Payment.find_one(payment_id)
        if not payment:
            print(f"[YOOKASSA] Платёж {payment_id} не найден в API")
            return False

        if payment.status != "succeeded":
            print(f"[YOOKASSA] Платёж {payment_id} имеет статус {payment.status}, не succeeded")
            return False

        try:
            if str(payment.amount.value) and payment.amount.currency != "RUB":
                print(f"[YOOKASSA] Платёж {payment_id} не в рублях")
                return False
        except Exception:
            pass

        return True
    except Exception as e:
        print(f"[YOOKASSA] Ошибка верификации платежа {payment_id}: {e}")
        return False
async def send_fns_receipt(amount: int, description: str, user_email: str = ""):
    """Отправляет чек в ФНС через API «Мой налог» (самозанятый)."""
    print(f"[NALOGO] Вызов send_fns_receipt: amount={amount}, desc={description[:50]}")
    inn = os.getenv("NALOGO_INN", "")
    password = os.getenv("NALOGO_PASSWORD", "")
    if not inn or not password:
        print("[NALOGO] ИНН или пароль не заданы — чек не отправлен")
        return None

    try:
        client = NalogClient(inn=inn, password=password)
        await client.create_new_access_token()
        receipt = await client.income().create(
            name=description,
            amount=amount,
        )
        link = client.receipt().print_url(receipt)
        print(f"[NALOGO] ✅ Чек создан: {link}")
        return link
    except Exception as e:
        print(f"[NALOGO] ❌ Ошибка создания чека: {e}")
        return None


# ================== ЧЁРНЫЙ СПИСОК ==================
def is_blacklisted(username: str = "", phone: str = "") -> bool:
    """Проверяет, есть ли username или phone в чёрном списке."""
    if username:
        u = username.lower().lstrip("@")
        for key, item in blacklist.items():
            if item.get("type") == "username":
                v = item.get("value", "").lower().lstrip("@")
                if v == u:
                    return True
    if phone:
        p = phone.replace(" ", "").replace("-", "").replace("+", "")
        for key, item in blacklist.items():
            if item.get("type") == "phone":
                v = item.get("value", "").replace(" ", "").replace("-", "").replace("+", "")
                if v == p:
                    return True
    return False


def add_to_blacklist(entry_type: str, value: str, reason: str = "") -> bool:
    """Добавляет запись в чёрный список."""
    value = value.strip()
    if not value:
        return False

    value_clean = value.lower().lstrip("@").replace(" ", "").replace("-", "").replace("+", "")
    for key, item in blacklist.items():
        existing = item.get("value", "").lower().lstrip("@").replace(" ", "").replace("-", "").replace("+", "")
        if item.get("type") == entry_type and existing == value_clean:
            return False

    key = f"{entry_type}:{value_clean}"
    blacklist[key] = {
        "type": entry_type,
        "value": value,
        "reason": reason,
        "added_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    try:
        supabase.table("blacklist").upsert({
            "type": entry_type,
            "value": value,
            "reason": reason,
            "added_at": blacklist[key]["added_at"],
        }).execute()
    except Exception as e:
        print(f"[SUPABASE] Ошибка добавления в ЧС: {e}")

    return True


def remove_from_blacklist(value: str) -> bool:
    """Удаляет запись из чёрного списка по value."""
    value_clean = value.lower().lstrip("@").replace(" ", "").replace("-", "").replace("+", "")
    found_key = None
    for key, item in blacklist.items():
        existing = item.get("value", "").lower().lstrip("@").replace(" ", "").replace("-", "").replace("+", "")
        if existing == value_clean:
            found_key = key
            break

    if not found_key:
        return False

    item = blacklist.pop(found_key)
    try:
        supabase.table("blacklist").delete().eq("value", item["value"]).execute()
    except Exception as e:
        print(f"[SUPABASE] Ошибка удаления из ЧС: {e}")
    return True


def get_blacklist_text() -> str:
    """Возвращает текстовое представление чёрного списка."""
    if not blacklist:
        return "📭 Чёрный список пуст."

    text = f"🚫 <b>Чёрный список</b> ({len(blacklist)} записей)\n\n"
    for key, item in blacklist.items():
        icon = "📱" if item.get("type") == "phone" else "🔗"
        reason = f" — {item.get('reason')}" if item.get("reason") else ""
        text += f"{icon} <code>{item.get('value')}</code>{reason}\n"
    return text


# ================== ЛОГ ПОПЫТОК РЕГИСТРАЦИИ ==================
def log_reg_attempt(
    user_id: int,
    username: str = "",
    first_name: str = "",
    phone: str = "",
    status: str = "started",
    step: str = "start",
    reason: str = "",
):
    """Логирует попытку регистрации."""
    attempt = {
        "user_id": user_id,
        "username": username or "",
        "first_name": first_name or "",
        "phone": phone or "",
        "status": status,
        "step": step,
        "reason": reason,
        "attempted_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    reg_attempts.append(attempt)

    try:
        supabase.table("reg_attempts").insert(attempt).execute()
    except Exception as e:
        print(f"[SUPABASE] Ошибка логирования попытки: {e}")


def get_reg_attempts_text(limit: int = 30) -> str:
    """Возвращает текст со свежими попытками регистрации."""
    if not reg_attempts:
        return "📭 Пока нет попыток регистрации."

    total = len(reg_attempts)
    blocked_user = sum(1 for a in reg_attempts if a.get("status") == "blocked_username")
    blocked_phone = sum(1 for a in reg_attempts if a.get("status") == "blocked_phone")
    success = sum(1 for a in reg_attempts if a.get("status") == "success")

    text = (
        f"📊 <b>Попытки регистрации</b>\n\n"
        f"Всего: <b>{total}</b>\n"
        f"✅ Успешных: <b>{success}</b>\n"
        f"🚫 Заблокирован username: <b>{blocked_user}</b>\n"
        f"🚫 Заблокирован телефон: <b>{blocked_phone}</b>\n\n"
        f"<b>Последние {min(limit, total)}:</b>\n\n"
    )

    for a in reg_attempts[-limit:][::-1]:
        status_icon = {
            "started": "🟡",
            "success": "✅",
            "blocked_username": "🚫",
            "blocked_phone": "🚫",
        }.get(a.get("status"), "❓")

        text += (
            f"{status_icon} <code>{a.get('user_id')}</code>"
            f" | @{a.get('username') or '—'}"
            f" | {a.get('first_name') or '—'}\n"
            f"   └ {a.get('status')} / шаг: {a.get('step')}"
            f" / {a.get('attempted_at')}\n"
        )
        if a.get("reason"):
            text += f"   └ причина: {a.get('reason')}\n"

    return text



# ================== УТИЛИТЫ ==================
def is_valid_email(email: str) -> bool:
    return re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email) is not None


def parse_event_date(text: str):
    """Парсит дату в свободной форме."""
    settings_parser = {
        "DATE_ORDER": "DMY",
        "PREFER_DATES_FROM": "future",
        "RETURN_AS_TIMEZONE_AWARE": False,
    }
    dt = dateparser.parse(text, languages=["ru", "en"], settings=settings_parser)
    if dt is None:
        return None, None

    months_ru = {
        1: "января", 2: "февраля", 3: "марта", 4: "апреля",
        5: "мая", 6: "июня", 7: "июля", 8: "августа",
        9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
    }
    date_text = f"{dt.day} {months_ru[dt.month]} {dt.year}, {dt.strftime('%H:%M')}"
    return dt, date_text


def sales_closed(event: dict) -> bool:
    start_dt = datetime.fromisoformat(event["start_dt"])
    stop_before = timedelta(hours=settings.get("sales_stop_hours", 1))
    return datetime.now() >= start_dt - stop_before


def is_event_archived(event: dict) -> bool:
    """Проверяет, пора ли событие в архив (6 часов после start_dt)."""
    try:
        start_dt = datetime.fromisoformat(event["start_dt"])
        archive_hours = settings.get("archive_after_hours", 6)
        return datetime.now() >= start_dt + timedelta(hours=archive_hours)
    except Exception:
        return False


def get_active_events() -> list:
    """Возвращает будущие (неархивные) события."""
    return [e for e in events if not is_event_archived(e)]


def get_archived_events() -> list:
    """Возвращает прошедшие события (архив)."""
    return [e for e in events if is_event_archived(e)]


def is_ticket_archived(ticket: dict) -> bool:
    """Проверяет, пора ли билет в архив (30 дней после used_at)."""
    if not ticket.get("used"):
        return False
    used_at_str = ticket.get("used_at") or ""
    if not used_at_str:
        return False
    try:
        used_at = datetime.strptime(used_at_str, "%Y-%m-%d %H:%M")
        days = settings.get("ticket_archive_days", 30)
        return datetime.now() >= used_at + timedelta(days=days)
    except Exception:
        return False


# ========== КОНЕЦ ЧАСТИ 3 ==========


def generate_ticket_code() -> str:
    return f"VIDA-{uuid.uuid4().hex[:4].upper()}-{uuid.uuid4().hex[:4].upper()}"


def next_order_id() -> str:
    return f"ORD-{uuid.uuid4().hex[:12].upper()}"


def make_qr_image(data: str) -> BufferedInputFile:
    """Генерирует PNG с QR-кодом для отправки гостю."""
    qr = qrcode.QRCode(version=1, box_size=10, border=2)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return BufferedInputFile(buf.getvalue(), filename=f"{data}.png")


def make_csv(headers, rows) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def format_user_info(user_id: int, u: dict) -> str:
    return (
        f"👤 <b>{u.get('name', '—')} {u.get('surname', '—')}</b>\n"
        f"🆔 <code>{user_id}</code>\n"
        f"📧 {u.get('email', '—') or '—'}\n"
        f"📱 {u.get('phone', '—')}\n"
        f"⚧ {u.get('gender', '—')}\n"
        f"🔗 @{u.get('username', '—') or '—'}\n"
        f"📅 {u.get('registered_at', '—')}\n"
    )


def get_short_event_label(event_date: str) -> str:
    """Возвращает короткую метку даты вида '15.11' из строки '15 ноября 2026, 22:00'."""
    if not event_date:
        return ""
    try:
        m = re.match(r"^(\d{1,2})\s+([а-яё]+)", event_date, re.IGNORECASE)
        if not m:
            return ""
        day = m.group(1).zfill(2)
        month_ru = m.group(2).lower()
        months_map = {
            "января": "01", "февраля": "02", "марта": "03",
            "апреля": "04", "мая": "05", "июня": "06",
            "июля": "07", "августа": "08", "сентября": "09",
            "октября": "10", "ноября": "11", "декабря": "12",
        }
        month_num = months_map.get(month_ru)
        if month_num:
            return f"{day}.{month_num}"
    except Exception:
        pass
    return ""


def get_ticket_price_from_orders(uid, event_id) -> int:
    """Ищет цену билета в orders по user_id и event_id."""
    for o in orders:
        if (str(o.get("user_id")) == str(uid)
                and str(o.get("event_id")) == str(event_id)):
            try:
                p = int(o.get("price", 0) or 0)
                if p > 0:
                    return p
            except (TypeError, ValueError):
                pass
            break
    return 0


def get_button_suffix(ticket: dict) -> str:
    """Возвращает суффикс для кнопки вида ' (15.11, 500 ₽)'."""
    date_part = get_short_event_label(ticket.get("event_date") or "")
    uid = ticket.get("user_id")
    ev_id = ticket.get("event_id")
    price = get_ticket_price_from_orders(uid, ev_id)
    price_part = f"{price} ₽" if price > 0 else ""

    parts = [p for p in [date_part, price_part] if p]
    if not parts:
        return ""
    return f" ({', '.join(parts)})"


def _ticket_type_label(ticket: dict) -> str:
    """Возвращает 'Парный' или 'Обычный'."""
    return "Парный" if ticket.get("type") == "pair" else "Обычный"


def get_ticket_status_icon(ticket: dict) -> str:
    """Возвращает иконку статуса билета: ⚪ сгорел, 🔴 погашен, 🟢 активен."""
    if ticket.get("auto_used"):
        return "⚪"
    if ticket.get("used"):
        return "🔴"
    return "🟢"




# ================== CHECKER-УТИЛИТЫ ==================

def find_checker_for_user(
    user_id: int = 0,
    username: str = "",
    phone: str = "",
) -> Optional[dict]:
    """
    Ищет checker'а по user_id, username или phone.
    Возвращает dict или None.
    """
    uname = (username or "").lstrip("@").lower()
    phone_clean = (phone or "").replace(" ", "").replace("-", "").replace("+", "")

    for chk_id, c in checkers.items():
        if not c.get("active", True):
            continue

        if user_id and c.get("user_id") and str(c["user_id"]) == str(user_id):
            return c

        if uname and c.get("username") and c["username"] == uname:
            return c

        if phone_clean:
            existing_phone = (c.get("phone") or "").replace(" ", "").replace("-", "").replace("+", "")
            if existing_phone and existing_phone == phone_clean:
                return c

    return None


def checker_can_access_event(checker: dict, event_id) -> bool:
    """Проверяет, есть ли у checker'а доступ к событию."""
    if not checker:
        return False
    if event_id is None:
        return False

    allowed = checker.get("event_id")
    if allowed is None:
        return True

    return str(allowed) == str(event_id)


def get_checker_events(checker: dict) -> list:
    """Возвращает список событий, доступных checker'у."""
    if not checker:
        return []

    if checker.get("event_id") is None:
        return list(events)

    ev_id = str(checker.get("event_id"))
    return [e for e in events if str(e.get("id")) == ev_id]


def activate_checker(chk_id: str, user_id: int):
    """Привязывает Telegram user_id к checker'у и сохраняет."""
    c = checkers.get(chk_id)
    if not c:
        return
    c["user_id"] = user_id
    c["active"] = True
    save_checker(chk_id)
    print(f"[CHECKER] Активирован {chk_id} → user_id={user_id}")


# ================== ЛОГ ДЕЙСТВИЙ С БИЛЕТАМИ ==================

def log_ticket_action(code: str, checker: dict, action: str = "used"):
    """Логирует действие checker'а с билетом."""
    log = {
        "ticket_code": code,
        "action": action,
        "checker_id": checker.get("id") if checker else None,
        "checker_name": (checker or {}).get("name", ""),
        "checker_user_id": (checker or {}).get("user_id"),
        "at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    ticket_logs.append(log)
    try:
        supabase.table("ticket_logs").insert(log).execute()
    except Exception as e:
        print(f"[SUPABASE] log_ticket_action ошибка: {e}")


# ================== EMAIL-ОТЧЁТ (УПРОЩЁННЫЙ) ==================

async def send_report_to_admin():
    """Отправляет отчёт главному админу прямо в Telegram."""
    total_revenue = sum(int(o.get("price", 0) or 0) for o in orders)
    tickets_active = sum(1 for t in tickets.values() if not t.get("used") and not t.get("auto_used"))
    tickets_used = sum(1 for t in tickets.values() if t.get("used") and not t.get("auto_used"))
    tickets_burned = sum(1 for t in tickets.values() if t.get("auto_used"))

    text = (
        f"📊 <b>Отчёт</b>\n\n"
        f"👥 Пользователей: <b>{len(users)}</b>\n"
        f"🎉 Событий: <b>{len(events)}</b>\n"
        f"🛒 Заказов: <b>{len(orders)}</b>\n"
        f"💰 Выручка: <b>{total_revenue} ₽</b>\n\n"
        f"🎫 Билетов активно: <b>{tickets_active}</b>\n"
        f"🔴 Погашено: <b>{tickets_used}</b>\n"
        f"⚪ Сгорело: <b>{tickets_burned}</b>\n"
        f"🚫 В ЧС: <b>{len(blacklist)}</b>\n"
    )
    try:
        await bot.send_message(ADMIN_ID, text, parse_mode="HTML")
    except Exception as e:
        print(f"[REPORT] Ошибка отправки: {e}")


# ================== АВТО-ПОГАШЕНИЕ БИЛЕТОВ ==================
async def auto_burn_expired_tickets():
    """Автоматически погашает билеты прошедших событий (для неявившихся)."""
    now = datetime.now()
    burned_codes = []
    affected_events = {}

    for code, ticket in tickets.items():
        if ticket.get("used"):
            continue
        if ticket.get("auto_used"):
            continue

        event_id = ticket.get("event_id")
        event = next((e for e in events if str(e["id"]) == str(event_id)), None)
        if not event:
            continue

        try:
            start_dt = datetime.fromisoformat(event["start_dt"])
        except Exception:
            continue

        burn_hours = settings.get("auto_burn_hours", 6)
        if now >= start_dt + timedelta(hours=burn_hours):
            burn_time = start_dt + timedelta(hours=burn_hours)
            ticket["used"] = True
            ticket["used_at"] = burn_time.strftime("%Y-%m-%d %H:%M")
            ticket["used_by"] = "Автоматически (не явился)"
            ticket["auto_used"] = True
            burned_codes.append(code)

            ev_title = event.get("title", "—")
            affected_events.setdefault(ev_title, 0)
            affected_events[ev_title] += 1

            print(f"[AUTO-BURN] Билет {code} авто-погашен (событие {event_id} прошло)")

            owner_id = ticket.get("user_id")
            if owner_id:
                try:
                    await bot.send_message(
                        int(owner_id),
                        f"⚪ <b>Билет погашен автоматически</b>\n\n"
                        f"🎫 Код: <code>{code}</code>\n"
                        f"🎉 {ticket.get('event_title', '—')}\n"
                        f"📅 {ticket.get('event_date', '—')}\n\n"
                        f"<i>Вы не пришли на мероприятие — билет сгорел.</i>",
                        parse_mode="HTML",
                    )
                except Exception:
                    pass

    if burned_codes:
        for code in burned_codes:
            save_ticket(code)
        print(f"[AUTO-BURN] Автоматически погашено билетов: {len(burned_codes)}")

        try:
            lines = [
                f"⚪ <b>Авто-погашение билетов</b>\n\n"
                f"Всего сгорело: <b>{len(burned_codes)}</b>\n\n"
            ]
            for title, cnt in affected_events.items():
                lines.append(f"🎉 {title}: <b>{cnt}</b>\n")

            await bot.send_message(ADMIN_ID, "".join(lines), parse_mode="HTML")
        except Exception:
            pass


async def _auto_burn_loop():
    """Фоновая задача: раз в час проверяет и сжигает билеты прошедших событий."""
    await asyncio.sleep(60)

    while True:
        try:
            await auto_burn_expired_tickets()
        except Exception as e:
            print(f"[AUTO-BURN] Ошибка: {e}")

        await asyncio.sleep(3600)


def start_auto_burn():
    """Запускает авто-сжигание билетов в фоне."""
    asyncio.create_task(_auto_burn_loop())


# ========== КОНЕЦ ЧАСТИ 4 ==========

# ================== СОСТОЯНИЯ ==================
class Reg(StatesGroup):
    phone = State()
    name = State()
    surname = State()
    email = State()
    gender = State()
    confirm = State()
    ticket_type = State()
    partner_input = State()
    payment = State()
    transfer_input = State() 

class Admin(StatesGroup):
    menu = State()
    add_title = State()
    add_date_text = State()
    add_place = State()
    add_description = State()
    add_price_single = State()
    add_price_pair = State()
    add_photo = State()
    checker_add_username = State()
    checker_add_phone = State()
    checker_add_name = State()
    checker_choose_event = State()
    blacklist_add_username = State()
    blacklist_add_phone = State()
    blacklist_reason = State()
    broadcast = State()
    find_ticket = State()


class CheckerStates(StatesGroup):
    menu = State()
    find_code = State()
    find_name = State()




# ================== /start (deep_link = invite) ==================
@dp.message(CommandStart(deep_link=True))
async def start_with_invite(message: Message, state: FSMContext, command: CommandObject):
    """Обработка /start с payload (invite_XXX)."""
    payload = command.args or ""
    if payload.startswith("invite_"):
        try:
            inviter_id = int(payload.replace("invite_", ""))
            if inviter_id != message.from_user.id:
                pending_invites[message.from_user.id] = inviter_id
                print(f"[INVITE] user={message.from_user.id} приглашён от {inviter_id}")
        except ValueError:
            pass

    await start(message, state)


# ================== /start ==================
@dp.message(CommandStart())
async def start(message: Message, state: FSMContext):
    await state.clear()
    user_id = message.from_user.id
    user_name = message.from_user.first_name or "друг"

    # ============ ГЛАВНЫЙ АДМИН ============
    if user_id == ADMIN_ID:
        greeting = (
            f"Привет, {user_name} 👋\n"
            "Вы вошли как <b>главный администратор</b>.\n\n"
            "Используйте панель ниже для управления ботом.\n\n"
            "Подпишитесь на соцсети:\n"
            "ВК (https://vk.com/vidabachata) • "
            "Telegram (https://t.me/+eTMoEG6V1883Nzli)"
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🛠 Админ-панель", callback_data="admin:back")],
            [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
            [InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets")],
            [InlineKeyboardButton(text="ВК", url="https://vk.com/vidabachata"),
             InlineKeyboardButton(text="Telegram", url="https://t.me/+eTMoEG6V1883Nzli")],
        ])
        await message.answer(greeting, reply_markup=kb, parse_mode="HTML")
        return

    # ============ CHECKER ============
    checker = find_checker_for_user(
        user_id=user_id,
        username=message.from_user.username or "",
    )

    if checker and not checker.get("user_id") and checker.get("phone"):
        await message.answer(
            f"👋 {user_name}, вы — админ на входе.\n\n"
            "Для активации доступа поделитесь контактом:",
            reply_markup=ReplyKeyboardMarkup(
                keyboard=[[KeyboardButton(text="📱 Поделиться контактом", request_contact=True)]],
                resize_keyboard=True, one_time_keyboard=True,
            ),
        )
        await state.update_data(checker_id=checker["id"])
        await state.set_state(CheckerStates.menu)
        return

    if checker:
        if not checker.get("user_id"):
            activate_checker(checker["id"], user_id)
        return await checker_start(message, state, checker)

    # ============ ЧС ============
    if is_blacklisted(username=message.from_user.username or ""):
        log_reg_attempt(
            user_id=user_id,
            username=message.from_user.username or "",
            first_name=user_name,
            status="blocked_username",
            step="start",
            reason="username в ЧС (/start)",
        )
        await message.answer("⚠️ Бот на техническом обслуживании. Попробуйте позже.")
        return

    # ============ ЗАРЕГИСТРИРОВАННЫЙ ============
    if user_id in users:
        greeting = (
            f"Привет, {user_name} 👋\n"
            "Здесь можно купить билет на вечеринки <b>Vida Bachata</b> 💃🕺\n\n"
            "Выберите мероприятие:\n\n"
            "Подпишитесь на соцсети:\n"
            "ВК (https://vk.com/vidabachata) • "
            "Telegram (https://t.me/+eTMoEG6V1883Nzli)"
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
            [InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets")],
            [InlineKeyboardButton(text="ВК", url="https://vk.com/vidabachata"),
             InlineKeyboardButton(text="Telegram", url="https://t.me/+eTMoEG6V1883Nzli")],
        ])
        await message.answer(greeting, reply_markup=kb, parse_mode="HTML")
        return

    # ============ НОВЫЙ ============
    greeting = (
        f"Привет, {user_name} 👋\n"
        "Здесь можно купить билет на вечеринки <b>Vida Bachata</b> 💃🕺\n\n"
        "Зарегистрируйтесь или посмотрите мероприятия.\n\n"
        "Подпишитесь на соцсети:\n"
        "ВК (https://vk.com/vidabachata) • "
        "Telegram (https://t.me/+eTMoEG6V1883Nzli)"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📝 Зарегистрироваться", callback_data="start_registration")],
        [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
        [InlineKeyboardButton(text="ВК", url="https://vk.com/vidabachata"),
         InlineKeyboardButton(text="Telegram", url="https://t.me/+eTMoEG6V1883Nzli")],
    ])
    await message.answer(greeting, reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "start_registration")
async def start_registration(call: CallbackQuery, state: FSMContext):
    if call.from_user.id == ADMIN_ID:
        await call.answer("Вы администратор, регистрация не нужна 👍", show_alert=True)
        return

    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if checker:
        await call.answer("Вы админ на входе 👍", show_alert=True)
        return

    if is_blacklisted(username=call.from_user.username or ""):
        log_reg_attempt(
            user_id=call.from_user.id,
            username=call.from_user.username or "",
            first_name=call.from_user.first_name or "",
            status="blocked_username",
            step="start_registration",
            reason="username в ЧС",
        )
        await call.answer("Доступ ограничен", show_alert=True)
        await call.message.answer("⚠️ Бот на техническом обслуживании. Попробуйте позже.")
        return

    if call.from_user.id in users:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
            [InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets")],
        ])
        await call.message.answer(
            "✅ Вы уже зарегистрированы!\n\nНажмите «Мероприятия».",
            reply_markup=kb,
        )
        return

    log_reg_attempt(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
        first_name=call.from_user.first_name or "",
        status="started",
        step="start",
    )

    print(f"[REG START] user={call.from_user.id} — запускаем регистрацию с телефона")

    kb = ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text="📱 Поделиться контактом", request_contact=True)]],
        resize_keyboard=True, one_time_keyboard=True,
    )
    await call.message.answer(
        "📱 Нажмите «Поделиться контактом», чтобы подтвердить ваш номер телефона:",
        reply_markup=kb,
    )
    await state.set_state(Reg.phone)


# ================== РЕГИСТРАЦИЯ ==================
@dp.message(Reg.phone, F.contact)
async def get_phone(message: Message, state: FSMContext):
    phone = message.contact.phone_number

    checker = find_checker_for_user(
        user_id=message.from_user.id,
        username=message.from_user.username or "",
        phone=phone,
    )
    if checker:
        activate_checker(checker["id"], message.from_user.id)
        await message.answer(
            "✅ Ваш номер подтверждён. Добро пожаловать!",
            reply_markup=ReplyKeyboardRemove(),
        )
        await state.clear()
        return await checker_start(message, state, checker)

    if is_blacklisted(phone=phone):
        log_reg_attempt(
            user_id=message.from_user.id,
            username=message.from_user.username or "",
            first_name=message.from_user.first_name or "",
            phone=phone,
            status="blocked_phone",
            step="phone",
            reason="телефон в ЧС",
        )
        await message.answer(
            "⚠️ Бот на техническом обслуживании. Попробуйте позже.",
            reply_markup=ReplyKeyboardRemove(),
        )
        await state.clear()
        return

    await state.update_data(phone=phone)
    await message.answer(
        "✅ Телефон получен!\n\n👤 Отправьте ваше Имя:",
        reply_markup=ReplyKeyboardRemove(),
    )
    await state.set_state(Reg.name)


@dp.message(Reg.phone)
async def wrong_phone(message: Message):
    await message.answer("Пожалуйста, нажмите кнопку «Поделиться контактом» 👇")


# ========== КОНЕЦ ЧАСТИ 5 ==========

@dp.message(Reg.name)
async def get_name(message: Message, state: FSMContext):
    name = message.text.strip()
    if len(name) < 2:
        await message.answer("Слишком короткое имя:")
        return
    await state.update_data(name=name)
    await message.answer("👤 Отправьте вашу Фамилию.")
    await state.set_state(Reg.surname)


@dp.message(Reg.surname)
async def get_surname(message: Message, state: FSMContext):
    surname = message.text.strip()
    if len(surname) < 2:
        await message.answer("Слишком короткая фамилия:")
        return
    await state.update_data(surname=surname)
    await message.answer(
        "📧 Отправьте ваш Email (для рассылки).\n"
        "Можно пропустить — отправьте <code>-</code>",
        parse_mode="HTML",
    )
    await state.set_state(Reg.email)


@dp.message(Reg.email)
async def get_email(message: Message, state: FSMContext):
    email_input = message.text.strip()

    if email_input == "-":
        await state.update_data(email="")
        kb = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="👨 Мужской", callback_data="gender:M"),
            InlineKeyboardButton(text="👩 Женский", callback_data="gender:F"),
        ]])
        await message.answer("📭 Email пропущен.\n\n⚧ Выберите ваш пол:", reply_markup=kb)
        await state.set_state(Reg.gender)
        return

    if not is_valid_email(email_input):
        await message.answer(
            "❌ Похоже, email введён неверно. Попробуйте ещё раз\n"
            "или отправьте <code>-</code>, чтобы пропустить:",
            parse_mode="HTML",
        )
        return

    await state.update_data(email=email_input)
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="👨 Мужской", callback_data="gender:M"),
        InlineKeyboardButton(text="👩 Женский", callback_data="gender:F"),
    ]])
    await message.answer(
        "✅ Email сохранён для рассылки.\n\n⚧ Выберите ваш пол:",
        reply_markup=kb,
    )
    await state.set_state(Reg.gender)


@dp.callback_query(F.data.startswith("gender:"), Reg.gender)
async def get_gender(call: CallbackQuery, state: FSMContext):
    gender = "Мужской" if call.data == "gender:M" else "Женский"
    await state.update_data(gender=gender)
    data = await state.get_data()

    email_line = data.get("email") or "—"
    text = (
        "📋 <b>Проверьте данные:</b>\n\n"
        f"📱 {data.get('phone', '—')}\n"
        f"👤 {data.get('name', '—')} {data.get('surname', '—')}\n"
        f"📧 {email_line}\n"
        f"⚧ {gender}"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Согласен", callback_data="confirm_yes")],
        [InlineKeyboardButton(text="🔄 Перезаполнить", callback_data="confirm_restart")],
    ])
    await call.message.edit_text(f"Пол: {gender}")
    await call.message.answer(text, parse_mode="HTML")
    await call.message.answer("Всё верно?", reply_markup=kb)
    await state.set_state(Reg.confirm)


@dp.callback_query(F.data == "confirm_yes", Reg.confirm)
async def confirm_yes(call: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    data["registered_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    data["username"] = call.from_user.username or ""
    users[call.from_user.id] = data
    if call.from_user.username:
        username_index[call.from_user.username.lstrip("@").strip().lower()] = call.from_user.id

    log_reg_attempt(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
        first_name=data.get("name", ""),
        phone=data.get("phone", ""),
        status="success",
        step="confirm",
    )

    save_user(call.from_user.id)

    inviter_id = pending_invites.pop(call.from_user.id, None)
    if inviter_id and inviter_id in users:
        try:
            await bot.send_message(
                inviter_id,
                f"🎉 Второй участник зарегистрировался!\n"
                f"👤 {data['name']} {data['surname']}"
            )
        except Exception:
            pass

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
        [InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets")],
    ])
    await call.message.edit_text(
        "✅ Регистрация завершена!\n\nНажмите «Мероприятия».",
        reply_markup=kb,
    )


@dp.callback_query(F.data == "confirm_restart", Reg.confirm)
async def confirm_restart(call: CallbackQuery, state: FSMContext):
    users.pop(call.from_user.id, None)
    if call.from_user.username:
        username_index.pop(call.from_user.username.lower(), None)
    await state.clear()
    kb = ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text="📱 Поделиться контактом", request_contact=True)]],
        resize_keyboard=True, one_time_keyboard=True,
    )
    await call.message.edit_text("🔄 Начинаем заново.")
    await call.message.answer(
        "📱 Нажмите «Поделиться контактом»:",
        reply_markup=kb,
    )
    await state.set_state(Reg.phone)




# ================== ПАНЕЛЬ CHECKER'А ==================
async def checker_start(message: Message, state: FSMContext, checker: dict):
    """Панель администратора на входе."""
    await state.clear()
    await state.set_state(CheckerStates.menu)

    name = checker.get("name") or message.from_user.first_name or "коллега"
    event_id = checker.get("event_id")

    if event_id is None:
        access_text = "🌐 <b>Все вечеринки</b>"
    else:
        ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
        if ev:
            access_text = f"🎯 <b>{ev['title']}</b>\n📅 {ev['date_text']}"
        else:
            access_text = f"🎯 Событие (id={event_id})"

    text = (
        f"👋 Привет, <b>{name}</b>!\n\n"
        f"🎫 Вы — <b>администратор на входе</b>.\n\n"
        f"Область доступа:\n{access_text}\n\n"
        "📥 Гости присылают билеты прямо сюда — "
        "с кнопками «Погасить» / «Отклонить».\n\n"
        "Выберите действие:"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔍 Найти билет по коду", callback_data="checker:find_code")],
        [InlineKeyboardButton(text="👤 Найти по имени/телефону", callback_data="checker:find_name")],
        [InlineKeyboardButton(text="📊 Статистика входа", callback_data="checker:stats")],
        [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
    ])

    await message.answer(text, reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "checker:menu")
async def checker_menu_callback(call: CallbackQuery, state: FSMContext):
    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return
    await state.clear()
    await checker_start(call.message, state, checker)
    await call.answer()


@dp.message(F.contact, CheckerStates.menu)
async def checker_phone_confirm(message: Message, state: FSMContext):
    """Обработка подтверждения телефона для checker'а."""
    phone = message.contact.phone_number

    checker = find_checker_for_user(
        user_id=message.from_user.id,
        username=message.from_user.username or "",
        phone=phone,
    )
    if not checker:
        await message.answer(
            "❌ Ваш номер не найден в списке админов на входе.\n"
            "Обратитесь к главному администратору.",
            reply_markup=ReplyKeyboardRemove(),
        )
        await state.clear()
        return

    activate_checker(checker["id"], message.from_user.id)
    await message.answer("✅ Номер подтверждён!", reply_markup=ReplyKeyboardRemove())
    await checker_start(message, state, checker)


@dp.message(CheckerStates.menu)
async def checker_any_message(message: Message, state: FSMContext):
    """Если checker пишет что-то вручную."""
    await message.answer(
        "Используйте кнопки ниже 👇\n"
        "Или отправьте /start для возврата в меню.",
    )


# ================== ПОИСК БИЛЕТА ПО КОДУ ==================
@dp.callback_query(F.data == "checker:find_code")
async def checker_find_code(call: CallbackQuery, state: FSMContext):
    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return

    await call.message.answer(
        "🔍 <b>Поиск билета</b>\n\n"
        "Отправьте <b>код билета</b> (например, <code>VIDA-XXXX-XXXX</code>).\n\n"
        "Код написан под QR-кодом в билете гостя.\n\n"
        "Для отмены — /start",
        parse_mode="HTML",
    )
    await state.set_state(CheckerStates.find_code)


@dp.message(CheckerStates.find_code)
async def checker_process_code(message: Message, state: FSMContext):
    checker = find_checker_for_user(
        user_id=message.from_user.id,
        username=message.from_user.username or "",
    )
    if not checker:
        await message.answer("⛔ Вы не админ на входе.")
        await state.clear()
        return

    code = message.text.strip().upper()

    ticket = tickets.get(code)
    if not ticket:
        await message.answer(
            f"❌ Билет с кодом <code>{code}</code> не найден.",
            parse_mode="HTML",
        )
        return

    if not checker_can_access_event(checker, ticket.get("event_id")):
        ev = next(
            (e for e in events if str(e["id"]) == str(ticket.get("event_id"))),
            None
        )
        ev_title = ev["title"] if ev else "другое событие"
        await message.answer(
            f"⛔ У вас нет доступа к этому событию.\n\n"
            f"Билет относится к: <b>{ev_title}</b>",
            parse_mode="HTML",
        )
        return

    await show_checker_ticket_card(message, code, ticket)


# ========== КОНЕЦ ЧАСТИ 6 ==========


# ================== ПОИСК ПО ИМЕНИ / ТЕЛЕФОНУ ==================
@dp.callback_query(F.data == "checker:find_name")
async def checker_find_name(call: CallbackQuery, state: FSMContext):
    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return

    await call.message.answer(
        "👤 <b>Поиск по имени или телефону</b>\n\n"
        "Отправьте имя или номер телефона гостя.\n\n"
        "Используется, если гость забыл код билета.\n\n"
        "Для отмены — /start",
        parse_mode="HTML",
    )
    await state.set_state(CheckerStates.find_name)


@dp.message(CheckerStates.find_name)
async def checker_process_name(message: Message, state: FSMContext):
    checker = find_checker_for_user(
        user_id=message.from_user.id,
        username=message.from_user.username or "",
    )
    if not checker:
        await message.answer("⛔ Вы не админ на входе.")
        await state.clear()
        return

    query = message.text.strip().lower()
    if len(query) < 2:
        await message.answer("Слишком короткий запрос. Введите минимум 2 символа.")
        return

    found = []
    for code, t in tickets.items():
        if not checker_can_access_event(checker, t.get("event_id")):
            continue

        holder = (t.get("holder") or "").lower()
        uid = t.get("user_id")

        match = query in holder

        if not match and uid:
            try:
                u = users.get(int(uid), {})
                full = f"{u.get('name', '')} {u.get('surname', '')}".lower()
                phone = (u.get("phone") or "").replace(" ", "").replace("-", "").replace("+", "")
                query_clean = query.replace(" ", "").replace("-", "").replace("+", "")
                if query in full or (query_clean and query_clean in phone):
                    match = True
            except (TypeError, ValueError):
                pass

        if match:
            found.append((code, t))

    if not found:
        await message.answer(
            f"❌ По запросу «{query}» ничего не найдено.\n\n"
            "Попробуйте другое имя или телефон.",
        )
        return

    text = f"🔍 Найдено: <b>{len(found)}</b>\n\n"
    buttons = []
    for code, t in found[:10]:
        status = get_ticket_status_icon(t)
        type_label = "Парный" if t.get("type") == "pair" else "Обычный"
        text += (
            f"{status} <code>{code}</code>\n"
            f"   👤 {t.get('holder', '—')}\n"
            f"   🎫 {type_label}\n"
            f"   🎉 {t.get('event_title', '—')}\n\n"
        )
        buttons.append([InlineKeyboardButton(
            text=f"Показать {code[-8:]}",
            callback_data=f"checker:show:{code}"
        )])

    buttons.append([InlineKeyboardButton(
        text="⬅️ В меню", callback_data="checker:menu"
    )])

    await message.answer(
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("checker:show:"))
async def checker_show_ticket(call: CallbackQuery):
    code = call.data.split(":", 2)[2]

    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if not checker_can_access_event(checker, ticket.get("event_id")):
        await call.answer("Нет доступа к этому событию", show_alert=True)
        return

    await call.answer()
    await show_checker_ticket_card(call.message, code, ticket)


# ================== КАРТОЧКА БИЛЕТА ДЛЯ CHECKER'А ==================
async def show_checker_ticket_card(message: Message, code: str, ticket: dict):
    """Показывает карточку билета с кнопкой погашения (для checker'а)."""
    if ticket.get("auto_used"):
        status = "⚪ сгорел (не явился)"
    elif ticket.get("used"):
        status = "🔴 использован"
    else:
        status = "🟢 активен"

    holder_name = ticket.get("holder", "—")
    holder_username = ""
    holder_phone = "—"

    uid = ticket.get("user_id")
    try:
        if uid:
            user = users.get(int(uid), {})
            holder_username = user.get("username") or ""
            holder_phone = user.get("phone") or "—"
    except (TypeError, ValueError):
        pass

    text = (
        f"🎫 <b>Билет найден</b>\n\n"
        f"Код: <code>{code}</code>\n"
        f"Статус: {status}\n\n"
        f"👤 {holder_name}\n"
    )
    if holder_username:
        text += f"🔗 @{holder_username}\n"
    text += f"📱 {holder_phone}\n\n"
    text += (
        f"🎫 {_ticket_type_label(ticket)}\n"
        f"🎉 {ticket.get('event_title', '—')}\n"
        f"📅 {ticket.get('event_date', '—')}\n"
        f"🕐 Выдан: {ticket.get('issued_at', '—')}\n"
    )

    kb_rows = []

    if not ticket.get("used") and not ticket.get("auto_used"):
        kb_rows.append([InlineKeyboardButton(
            text="✅ Погасить билет",
            callback_data=f"checker:use:{code}"
        )])
    elif ticket.get("auto_used"):
        kb_rows.append([InlineKeyboardButton(
            text="⚪ Билет уже недействителен",
            callback_data="checker:noop"
        )])

    kb_rows.append([InlineKeyboardButton(
        text="🔄 Обновить",
        callback_data=f"checker:show:{code}"
    )])
    kb_rows.append([InlineKeyboardButton(
        text="⬅️ В меню", callback_data="checker:menu"
    )])
    kb_rows.append([InlineKeyboardButton(
        text="🏠 На главную", callback_data="main_menu"
    )])

    await message.answer(
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows),
        parse_mode="HTML",
    )


@dp.callback_query(F.data == "checker:noop")
async def checker_noop(call: CallbackQuery):
    await call.answer("Этот билет уже недействителен", show_alert=True)




# ================== КАРТОЧКА БИЛЕТА ДЛЯ ГЛАВНОГО АДМИНА ==================
async def show_super_admin_ticket_card(message: Message, code: str, ticket: dict):
    """Карточка билета для главного админа."""
    if ticket.get("auto_used"):
        status = "⚪ сгорел (не явился)"
    elif ticket.get("used"):
        status = "🔴 использован"
    else:
        status = "🟢 активен"

    holder_name = ticket.get("holder", "—")
    holder_username = ""
    holder_phone = "—"

    uid = ticket.get("user_id")
    try:
        if uid:
            user = users.get(int(uid), {})
            holder_username = user.get("username") or ""
            holder_phone = user.get("phone") or "—"
    except (TypeError, ValueError):
        pass

    text = (
        f"🎫 <b>Билет найден</b> <i>(главный админ)</i>\n\n"
        f"Код: <code>{code}</code>\n"
        f"Статус: {status}\n\n"
        f"👤 {holder_name}\n"
    )
    if holder_username:
        text += f"🔗 @{holder_username}\n"
    text += f"📱 {holder_phone}\n\n"
    text += (
        f"🎫 {_ticket_type_label(ticket)}\n"
        f"🎉 {ticket.get('event_title', '—')}\n"
        f"📅 {ticket.get('event_date', '—')}\n"
        f"🕐 Выдан: {ticket.get('issued_at', '—')}\n"
    )

    if ticket.get("used"):
        text += (
            f"✅ Погашен: {ticket.get('used_at', '—')}\n"
            f"👮 Кем: {ticket.get('used_by', '—')}\n"
        )

    kb_rows = []

    if not ticket.get("used") and not ticket.get("auto_used"):
        kb_rows.append([InlineKeyboardButton(
            text="✅ Погасить билет",
            callback_data=f"admin:use_qr:{code}"
        )])
    else:
        kb_rows.append([InlineKeyboardButton(
            text="♻️ Восстановить билет",
            callback_data=f"admin:restore_ticket:{code}"
        )])

    kb_rows.append([InlineKeyboardButton(
        text="⬅️ В админ-панель", callback_data="admin:back"
    )])
    kb_rows.append([InlineKeyboardButton(
        text="🏠 На главную", callback_data="main_menu"
    )])

    await message.answer(
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows),
        parse_mode="HTML",
    )


# ================== ПОГАШЕНИЕ БИЛЕТА CHECKER'ОМ ==================
@dp.callback_query(F.data.startswith("checker:use:"))
async def checker_use_ticket(call: CallbackQuery):
    code = call.data.split(":", 2)[2]

    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if not checker_can_access_event(checker, ticket.get("event_id")):
        await call.answer("Нет доступа к этому событию", show_alert=True)
        return

    if ticket.get("auto_used"):
        await call.answer("Этот билет уже недействителен", show_alert=True)
        return

    if ticket.get("used"):
        await call.answer("Билет уже погашен", show_alert=True)
        return

    ticket["used"] = True
    ticket["used_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ticket["used_by"] = checker.get("name") or call.from_user.full_name
    save_ticket(code)

    log_ticket_action(code, checker, action="used")

    try:
        checker_name = checker.get("name") or call.from_user.full_name
        await bot.send_message(
            ADMIN_ID,
            f"🎫 <b>Билет погашен</b>\n\n"
            f"Код: <code>{code}</code>\n"
            f"🎉 {ticket.get('event_title', '—')}\n"
            f"👤 {ticket.get('holder', '—')}\n\n"
            f"Погасил: <b>{checker_name}</b>",
            parse_mode="HTML",
        )
    except Exception:
        pass

    try:
        owner_id = ticket.get("user_id")
        if owner_id:
            await bot.send_message(
                owner_id,
                f"🔴 <b>Билет погашен!</b>\n\n"
                f"🎫 Код: <code>{code}</code>\n"
                f"🎉 {ticket.get('event_title', '—')}\n"
                f"📅 {ticket.get('event_date', '—')}\n"
                f"👤 {ticket.get('holder', '—')}\n\n"
                f"Хорошего вечера! 💃🕺",
                parse_mode="HTML",
            )
    except Exception as e:
        print(f"[NOTIFY] Ошибка уведомления гостя: {e}")

        # Собираем данные гостя
    holder_user = {}
    try:
        if ticket.get("user_id"):
            holder_user = users.get(int(ticket.get("user_id")), {})
    except (TypeError, ValueError):
        pass
    holder_username = holder_user.get("username") or ""
    holder_phone = holder_user.get("phone") or "—"

    checker_name = checker.get("name") or call.from_user.full_name
    kb_ok = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ В меню", callback_data="checker:menu")],
        [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
    ])

    text_ok = (
        f"🔴 <b>Билет погашен</b>\n\n"
        f"Код: <code>{code}</code>\n"
        f"👤 {ticket.get('holder', '—')}\n"
    )
    if holder_username:
        text_ok += f"🔗 @{holder_username}\n"
    text_ok += f"📱 {holder_phone}\n\n"
    text_ok += (
        f"🎫 {_ticket_type_label(ticket)}\n"
        f"🎉 {ticket.get('event_title', '—')}\n"
        f"📅 {ticket.get('event_date', '—')}\n\n"
        f"👮 Погасил: {checker_name}\n"
        f"🕐 {ticket.get('used_at', '—')}"
    )

    await call.message.edit_text(
        text_ok,
        reply_markup=kb_ok,
        parse_mode="HTML",
    )


# ================== ПОГАШЕНИЕ ГЛАВНЫМ АДМИНОМ ==================
@dp.callback_query(F.data.startswith("admin:use_qr:"))
async def admin_use_ticket_qr(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    code = call.data.split(":", 2)[2]
    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if ticket.get("auto_used"):
        await call.answer("Этот билет уже недействителен", show_alert=True)
        return

    if ticket.get("used"):
        await call.answer("Уже погашен", show_alert=True)
        return

    ticket["used"] = True
    ticket["used_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ticket["used_by"] = "Главный админ"
    save_ticket(code)

    try:
        owner_id = ticket.get("user_id")
        if owner_id:
            await bot.send_message(
                int(owner_id),
                f"✅ <b>Ваш билет погашен</b>\n\n"
                f"Код: <code>{code}</code>\n"
                f"Администратор: главный админ\n"
                f"Время: {ticket['used_at']}",
                parse_mode="HTML",
            )
    except Exception:
        pass

    await call.answer("✅ Погашен")
    await show_super_admin_ticket_card(call.message, code, ticket)


# ================== ВОССТАНОВЛЕНИЕ БИЛЕТА ==================
@dp.callback_query(F.data.startswith("admin:restore_ticket:"))
async def admin_restore_ticket(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    code = call.data.split(":", 2)[2]
    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if not ticket.get("used") and not ticket.get("auto_used"):
        await call.answer("Билет уже активен", show_alert=True)
        return

    ticket["used"] = False
    ticket["auto_used"] = False
    ticket["used_at"] = ""
    ticket["used_by"] = ""
    save_ticket(code)

    print(f"[RESTORE] Билет {code} восстановлен админом")

    try:
        owner_id = ticket.get("user_id")
        if owner_id:
            await bot.send_message(
                int(owner_id),
                f"♻️ <b>Ваш билет восстановлен!</b>\n\n"
                f"🎫 Код: <code>{code}</code>\n"
                f"🎉 {ticket.get('event_title', '—')}\n"
                f"📅 {ticket.get('event_date', '—')}\n"
                f"👤 {ticket.get('holder', '—')}\n\n"
                f"Покажите QR на входе.",
                parse_mode="HTML",
            )
    except Exception as e:
        print(f"[RESTORE] Ошибка уведомления гостя: {e}")

    await call.answer("♻️ Билет восстановлен")
    await show_super_admin_ticket_card(call.message, code, ticket)


# ================== ОТКЛОНЕНИЕ БИЛЕТА CHECKER'ОМ ==================
@dp.callback_query(F.data.startswith("checker:reject:"))
async def checker_reject_ticket(call: CallbackQuery):
    """Checker отклоняет присланный билет."""
    code = call.data.split(":", 2)[2]

    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    log_ticket_action(code, checker, action="rejected")

    try:
        owner_id = ticket.get("user_id")
        if owner_id:
            await bot.send_message(
                owner_id,
                f"❌ <b>Проверяющий отклонил билет</b>\n\n"
                f"Код: <code>{code}</code>\n\n"
                f"Обратитесь к организатору.",
                parse_mode="HTML",
            )
    except Exception:
        pass

    checker_name = checker.get("name") or call.from_user.full_name
    kb_rej = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ В меню", callback_data="checker:menu")],
        [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
    ])
    await call.message.edit_text(
        f"❌ <b>Билет отклонён</b>\n\n"
        f"Код: <code>{code}</code>\n"
        f"👤 {ticket.get('holder', '—')}\n\n"
        f"Отклонил: {checker_name}",
        reply_markup=kb_rej,
        parse_mode="HTML",
    )


# ================== СТАТИСТИКА CHECKER'А ==================
@dp.callback_query(F.data == "checker:stats")
async def checker_stats(call: CallbackQuery):
    checker = find_checker_for_user(
        user_id=call.from_user.id,
        username=call.from_user.username or "",
    )
    if not checker:
        await call.answer("Вы не админ на входе", show_alert=True)
        return

    accessible_events = get_checker_events(checker)
    event_ids = {str(e["id"]) for e in accessible_events}

    total = 0
    used = 0
    auto_used = 0
    for code, t in tickets.items():
        if str(t.get("event_id")) not in event_ids:
            continue
        total += 1
        if t.get("auto_used"):
            auto_used += 1
        elif t.get("used"):
            used += 1

    if checker.get("event_id") is None:
        header = "🌐 <b>Все вечеринки</b>"
    else:
        ev = next((e for e in events if str(e["id"]) == str(checker.get("event_id"))), None)
        header = f"🎯 <b>{ev['title'] if ev else 'Событие'}</b>"

    text = (
        f"📊 <b>Статистика входа</b>\n\n"
        f"{header}\n\n"
        f"🎫 Всего билетов: <b>{total}</b>\n"
        f"🔴 Погашено: <b>{used}</b>\n"
        f"⚪ Сгорело: <b>{auto_used}</b>\n"
        f"🟢 Осталось: <b>{total - used - auto_used}</b>\n"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="checker:stats")],
        [InlineKeyboardButton(text="⬅️ В меню", callback_data="checker:menu")],
        [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
    ])

    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


# ========== КОНЕЦ ЧАСТИ 7 ==========


# ================== АДМИН ==================
@dp.message(Command("admin"))
async def admin_start(message: Message, state: FSMContext):
    if message.from_user.id != ADMIN_ID:
        await message.answer("⛔ Только для админа.")
        return
    await state.clear()
    await show_admin_menu(message, state)


async def show_admin_menu(message: Message, state: FSMContext):
    active_cnt = len(get_active_events())
    archived_cnt = len(get_archived_events())
    auto_burned_cnt = sum(1 for t in tickets.values() if t.get("auto_used"))

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Создать событие", callback_data="admin:add")],
        [InlineKeyboardButton(text="🗑 Удалить событие", callback_data="admin:delete")],
        [InlineKeyboardButton(text="📋 Список событий", callback_data="admin:list")],
        [InlineKeyboardButton(text=f"📦 Архив событий ({archived_cnt})", callback_data="admin:archive")],
        [InlineKeyboardButton(text="🎉 Анонсы (просмотр)", callback_data="show_events")],
        [InlineKeyboardButton(text="🎫 Админы на входе", callback_data="admin:checkers")],
        [InlineKeyboardButton(text="👥 Пользователи", callback_data="admin:users:0")],
        [InlineKeyboardButton(text="🚫 Чёрный список", callback_data="admin:blacklist")],
        [InlineKeyboardButton(text="📊 Попытки регистрации", callback_data="admin:reg_attempts")],
        [InlineKeyboardButton(text="📈 Статистика", callback_data="admin:stats")],
        [InlineKeyboardButton(text="📢 Рассылка", callback_data="admin:broadcast")],
        [InlineKeyboardButton(text=f"🎟 Все билеты", callback_data="admin:tickets")],
        [InlineKeyboardButton(text="📋 История передач", callback_data="admin:transfers")],
        [InlineKeyboardButton(text=f"📦 Архив билетов ({auto_burned_cnt})", callback_data="admin:archived_tickets")],
        [InlineKeyboardButton(text="📤 Выгрузка билетов", callback_data="admin:export")],
        [InlineKeyboardButton(text="🎫 Найти билет", callback_data="admin:find_ticket")],
        [InlineKeyboardButton(text="📧 Отправить отчёт", callback_data="admin:report")],
    ])
    await message.answer(
        f"🛠 <b>Админ-панель</b>\n\n"
        f"🟢 Активных событий: {active_cnt} | 📦 В архиве: {archived_cnt}\n"
        f"Заказов: {len(orders)} | Юзеров: {len(users)}\n"
        f"ЧС: {len(blacklist)} | Попыток: {len(reg_attempts)} | Checkers: {len(checkers)}",
        reply_markup=kb, parse_mode="HTML",
    )
    await state.set_state(Admin.menu)


@dp.callback_query(F.data == "admin:back")
async def admin_back(call: CallbackQuery, state: FSMContext):
    await state.clear()
    await show_admin_menu(call.message, state)


@dp.callback_query(F.data == "admin:report")
async def admin_report(call: CallbackQuery):
    await call.answer("Отправляю...")
    await send_report_to_admin()
    await call.message.answer("📤 Отчёт отправлен в чат.")


# ================== АРХИВ СОБЫТИЙ ==================
@dp.callback_query(F.data == "admin:archive")
async def admin_archive(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    archived = get_archived_events()
    if not archived:
        await call.message.edit_text(
            "📭 Архив событий пуст.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")]
            ]),
        )
        return

    archived_sorted = sorted(archived, key=lambda e: e.get("start_dt", ""), reverse=True)

    text = f"📦 <b>Архив событий</b> ({len(archived_sorted)})\n\n"
    for e in archived_sorted[:10]:
        ev_tickets = [t for t in tickets.values() if str(t.get("event_id")) == str(e["id"])]
        total = len(ev_tickets)
        used = sum(1 for t in ev_tickets if t.get("used") and not t.get("auto_used"))
        auto_used = sum(1 for t in ev_tickets if t.get("auto_used"))
        revenue = sum(
            int(o.get("price", 0))
            for o in orders
            if str(o.get("event_id")) == str(e["id"])
        )

        text += (
            f"📦 <b>{e['title']}</b>\n"
            f"   📅 {e['date_text']}\n"
            f"   📍 {e.get('place', '—')}\n"
            f"   🎫 Билетов: {total}\n"
            f"   🔴 Погашено: {used}\n"
            f"   ⚪ Сгорело: {auto_used}\n"
            f"   💰 Выручка: {revenue} ₽\n"
            f"   🆔 <code>{e['id']}</code>\n\n"
        )

    if len(archived_sorted) > 10:
        text += f"... и ещё {len(archived_sorted) - 10}\n"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🗑 Удалить из архива", callback_data="admin:archive_delete")],
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:archive")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:archive_delete")
async def admin_archive_delete(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    archived = get_archived_events()
    if not archived:
        await call.answer("Архив пуст", show_alert=True)
        return

    archived_sorted = sorted(archived, key=lambda e: e.get("start_dt", ""), reverse=True)
    buttons = []
    for e in archived_sorted:
        buttons.append([InlineKeyboardButton(
            text=f"🗑 📦 {e['title'][:30]} — {e['date_text']}",
            callback_data=f"admin:del_arch:{e['id']}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:archive")])

    await call.message.edit_text(
        "🗑 <b>Удалить событие из архива?</b>\n\n"
        "Это действие удалит событие полностью (билеты останутся в базе).",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


# ================== РАЗДЕЛ «АДМИНЫ НА ВХОДЕ» ==================
@dp.callback_query(F.data == "admin:checkers")
async def admin_checkers_menu(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    total = len(checkers)
    all_events_cnt = sum(1 for c in checkers.values() if c.get("event_id") is None)
    specific_cnt = total - all_events_cnt
    active_cnt = sum(1 for c in checkers.values() if c.get("user_id"))

    text = (
        f"🎫 <b>Админы на входе</b>\n\n"
        f"Всего: <b>{total}</b>\n"
        f"🌐 На все вечеринки: <b>{all_events_cnt}</b>\n"
        f"🎯 На конкретные: <b>{specific_cnt}</b>\n"
        f"🟢 Активировано: <b>{active_cnt}</b>\n\n"
        f"Что делаем?"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Добавить по @username", callback_data="admin:chk:add_username")],
        [InlineKeyboardButton(text="📱 Добавить по телефону", callback_data="admin:chk:add_phone")],
        [InlineKeyboardButton(text="📋 Список всех", callback_data="admin:chk:list")],
        [InlineKeyboardButton(text="🗑 Удалить", callback_data="admin:chk:remove")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])

    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:chk:add_username")
async def admin_chk_add_username(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    await state.update_data(chk_type="username")
    await call.message.edit_text(
        "➕ <b>Добавление админа на входе</b>\n\n"
        "Шаг 1/3. Отправьте <b>@username</b> (например, @kolya_vhod):\n\n"
        "Для отмены — /admin",
        parse_mode="HTML",
    )
    await state.set_state(Admin.checker_add_username)


@dp.message(Admin.checker_add_username)
async def admin_chk_save_username(message: Message, state: FSMContext):
    if message.text == "/admin":
        await state.clear()
        await show_admin_menu(message, state)
        return

    value = message.text.strip()
    if not value.startswith("@"):
        value = "@" + value

    if len(value) < 3:
        await message.answer("❌ Слишком короткий username. Попробуйте ещё раз:")
        return

    value_clean = value.lstrip("@").lower()
    for chk_id, c in checkers.items():
        if c.get("username") == value_clean:
            await message.answer(f"⚠️ Checker с username {value} уже существует.")
            await state.clear()
            await show_admin_menu(message, state)
            return

    await state.update_data(chk_value=value)
    await state.set_state(Admin.checker_add_name)
    await message.answer(
        "Шаг 2/3. Отправьте <b>имя</b> checker'а (для отображения):\n\n"
        "Например: <code>Коля (вход)</code>\n\n"
        "Или отправьте <code>-</code>, чтобы оставить пустым:",
        parse_mode="HTML",
    )


@dp.callback_query(F.data == "admin:chk:add_phone")
async def admin_chk_add_phone(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    await state.update_data(chk_type="phone")
    await call.message.edit_text(
        "➕ <b>Добавление админа на входе</b>\n\n"
        "Шаг 1/3. Отправьте <b>номер телефона</b> (например, +79991234567):\n\n"
        "Для отмены — /admin",
        parse_mode="HTML",
    )
    await state.set_state(Admin.checker_add_phone)


@dp.message(Admin.checker_add_phone)
async def admin_chk_save_phone(message: Message, state: FSMContext):
    if message.text == "/admin":
        await state.clear()
        await show_admin_menu(message, state)
        return

    value = message.text.strip()
    value_clean = value.replace(" ", "").replace("-", "").replace("+", "")

    if len(value_clean) < 10:
        await message.answer("❌ Слишком короткий номер. Попробуйте ещё раз:")
        return

    for chk_id, c in checkers.items():
        existing = (c.get("phone") or "").replace(" ", "").replace("-", "").replace("+", "")
        if existing == value_clean:
            await message.answer(f"⚠️ Checker с телефоном {value} уже существует.")
            await state.clear()
            await show_admin_menu(message, state)
            return

    await state.update_data(chk_value=value)
    await state.set_state(Admin.checker_add_name)
    await message.answer(
        "Шаг 2/3. Отправьте <b>имя</b> checker'а (для отображения):\n\n"
        "Например: <code>Коля (вход)</code>\n\n"
        "Или отправьте <code>-</code>, чтобы оставить пустым:",
        parse_mode="HTML",
    )


@dp.message(Admin.checker_add_name)
async def admin_chk_save_name(message: Message, state: FSMContext):
    if message.text == "/admin":
        await state.clear()
        await show_admin_menu(message, state)
        return

    name = message.text.strip()
    if name == "-":
        name = ""

    await state.update_data(chk_name=name)

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🌐 Все вечеринки", callback_data="admin:chk:access:all")],
        [InlineKeyboardButton(text="🎯 Конкретная вечеринка", callback_data="admin:chk:access:specific")],
        [InlineKeyboardButton(text="❌ Отмена", callback_data="admin:checkers")],
    ])
    await message.answer(
        "Шаг 3/3. Выберите <b>область доступа</b>:",
        reply_markup=kb, parse_mode="HTML",
    )
    await state.set_state(Admin.checker_choose_event)


@dp.callback_query(F.data == "admin:chk:access:all")
async def admin_chk_access_all(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    data = await state.get_data()
    await finish_add_checker(call, state, event_id=None, data=data)


@dp.callback_query(F.data == "admin:chk:access:specific")
async def admin_chk_access_specific(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    active_events = get_active_events()
    if not active_events:
        await call.answer("Нет активных событий. Сначала создайте событие.", show_alert=True)
        return

    buttons = []
    for e in active_events:
        buttons.append([InlineKeyboardButton(
            text=f"🎉 {e['title']} — {e['date_text']}",
            callback_data=f"admin:chk:ev:{e['id']}"
        )])
    buttons.append([InlineKeyboardButton(text="❌ Отмена", callback_data="admin:checkers")])

    await call.message.edit_text(
        "🎯 Выберите <b>событие</b>, к которому даёте доступ:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("admin:chk:ev:"))
async def admin_chk_choose_event(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    event_id = call.data.split(":", 3)[3]
    data = await state.get_data()
    await finish_add_checker(call, state, event_id=event_id, data=data)


async def finish_add_checker(call: CallbackQuery, state: FSMContext, event_id: Optional[str], data: dict):
    """Завершает добавление checker'а."""
    chk_type = data.get("chk_type", "username")
    chk_value = data.get("chk_value", "")
    chk_name = data.get("chk_name", "")

    if not chk_value:
        await call.answer("Ошибка: нет значения", show_alert=True)
        await state.clear()
        return

    chk_id = f"chk_{uuid.uuid4().hex[:8]}"

    new_checker = {
        "id": chk_id,
        "user_id": None,
        "username": chk_value.lstrip("@").lower() if chk_type == "username" else "",
        "phone": chk_value if chk_type == "phone" else "",
        "name": chk_name,
        "event_id": event_id,
        "added_by": ADMIN_ID,
        "added_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "active": True,
    }

    checkers[chk_id] = new_checker
    save_checker(chk_id)

    if event_id is None:
        access_text = "🌐 Все вечеринки"
    else:
        ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
        access_text = f"🎯 {ev['title']} — {ev['date_text']}" if ev else f"🎯 Событие {event_id}"

    if chk_type == "username":
        value_text = f"🔗 @{new_checker['username']}"
    else:
        value_text = f"📱 {new_checker['phone']}"

    name_text = new_checker['name'] or "—"

    await state.clear()

    await call.message.edit_text(
        f"✅ <b>Админ на входе добавлен!</b>\n\n"
        f"👤 {name_text}\n"
        f"{value_text}\n"
        f"Область: {access_text}\n\n"
        f"<i>Он активируется, когда напишет /start боту "
        f"или поделится контактом (если добавлен по телефону).</i>",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="⬅️ К списку админов", callback_data="admin:checkers")
        ]]),
    )


# ========== КОНЕЦ ЧАСТИ 8 ==========


@dp.callback_query(F.data == "admin:chk:list")
async def admin_chk_list(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    if not checkers:
        await call.message.edit_text(
            "📭 Нет админов на входе.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:checkers")
            ]]),
        )
        return

    text = f"📋 <b>Админы на входе</b> ({len(checkers)})\n\n"
    for chk_id, c in checkers.items():
        status_icon = "🟢" if c.get("user_id") else "🟡"

        name = c.get("name") or "—"
        if c.get("username"):
            value = f"🔗 @{c['username']}"
        elif c.get("phone"):
            value = f"📱 {c['phone']}"
        else:
            value = "—"

        if c.get("event_id") is None:
            access = "🌐 Все вечеринки"
        else:
            ev = next((e for e in events if str(e["id"]) == str(c["event_id"])), None)
            access = f"🎯 {ev['title']}" if ev else f"🎯 {c['event_id']}"

        text += f"{status_icon} <b>{name}</b>\n   {value}\n   {access}\n\n"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:chk:list")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:checkers")],
    ])
    await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:chk:remove")
async def admin_chk_remove_start(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    if not checkers:
        await call.answer("Список пуст", show_alert=True)
        return

    buttons = []
    for chk_id, c in checkers.items():
        name = c.get("name") or "—"
        if c.get("username"):
            value = f"@{c['username']}"
        elif c.get("phone"):
            value = c["phone"]
        else:
            value = "—"
        buttons.append([InlineKeyboardButton(
            text=f"🗑 {name} — {value}",
            callback_data=f"admin:chk:del:{chk_id}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:checkers")])

    await call.message.edit_text(
        "🗑 Выберите админа для удаления:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
    )


@dp.callback_query(F.data.startswith("admin:chk:del:"))
async def admin_chk_remove_confirm(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для главного админа", show_alert=True)
        return

    chk_id = call.data.split(":", 3)[3]
    if chk_id not in checkers:
        await call.answer("Не найден", show_alert=True)
        return

    # ВАЖНО: сначала удаляем из БД (пока checker ещё в памяти),
    # потом — из памяти
    name = checkers[chk_id].get("name") or "—"
    delete_checker(chk_id)
    checkers.pop(chk_id, None)

    await call.answer(f"Удалён: {name}")
    await admin_checkers_menu(call, state)




# ================== СОЗДАНИЕ СОБЫТИЙ ==================
@dp.callback_query(F.data == "admin:add")
async def admin_add_start(call: CallbackQuery, state: FSMContext):
    await state.clear()
    await state.set_state(Admin.add_title)
    await call.message.edit_text(
        "➕ <b>Новое событие</b>\n\n"
        "Шаг 1/7. Введите <b>название</b> события:",
        parse_mode="HTML",
    )


@dp.message(Admin.add_title)
async def add_title(message: Message, state: FSMContext):
    await state.update_data(title=message.text.strip())
    await state.set_state(Admin.add_date_text)
    await message.answer(
        "Шаг 2/7. Введите <b>дату и время</b> события в любой удобной форме:\n\n"
        "Например:\n"
        "• <code>15.11.2026 22:00</code>\n"
        "• <code>15/11/2026 22:00</code>\n"
        "• <code>15 ноября 2026, 22:00</code>\n"
        "• <code>15 ноября 22:00</code>",
        parse_mode="HTML",
    )


@dp.message(Admin.add_date_text)
async def add_date_text(message: Message, state: FSMContext):
    text = message.text.strip()
    dt, date_text = parse_event_date(text)

    if dt is None:
        await message.answer(
            "❌ Не могу распознать дату.\n\n"
            "Попробуй один из форматов:\n"
            "• <code>15.11.2026 22:00</code>\n"
            "• <code>15/11/2026 22:00</code>\n"
            "• <code>15 ноября 2026, 22:00</code>\n"
            "• <code>15 ноября 22:00</code>",
            parse_mode="HTML",
        )
        return

    if dt < datetime.now():
        await message.answer("❌ Дата уже в прошлом. Введи будущую дату.")
        return

    await state.update_data(start_dt=dt.isoformat(), date_text=date_text)
    await state.set_state(Admin.add_place)
    await message.answer(
        f"✅ Распознано: <b>{date_text}</b>\n\n"
        f"Шаг 3/7. Место проведения:",
        parse_mode="HTML",
    )


@dp.message(Admin.add_place)
async def add_place(message: Message, state: FSMContext):
    await state.update_data(place=message.text.strip())
    await state.set_state(Admin.add_description)
    await message.answer("Шаг 4/7. Описание:")


@dp.message(Admin.add_description)
async def add_description(message: Message, state: FSMContext):
    await state.update_data(description=message.text.strip())
    await state.set_state(Admin.add_price_single)
    await message.answer("Шаг 5/7. Цена обычного (число):")


@dp.message(Admin.add_price_single)
async def add_price_single(message: Message, state: FSMContext):
    try:
        price = int(message.text.strip())
        if price <= 0:
            raise ValueError
    except ValueError:
        await message.answer("❌ Число > 0")
        return
    await state.update_data(price_single=price)
    await state.set_state(Admin.add_price_pair)
    await message.answer("Шаг 6/7. Цена парного (0 если нет):")


@dp.message(Admin.add_price_pair)
async def add_price_pair(message: Message, state: FSMContext):
    try:
        price = int(message.text.strip())
        if price < 0:
            raise ValueError
    except ValueError:
        await message.answer("❌ Число ≥ 0")
        return
    await state.update_data(price_pair=price)
    await state.set_state(Admin.add_photo)
    await message.answer("Шаг 7/7. Фото-обложка или <code>-</code>:", parse_mode="HTML")


@dp.message(Admin.add_photo, F.photo)
async def add_photo(message: Message, state: FSMContext):
    await state.update_data(photo_id=message.photo[-1].file_id)
    await finish_add_event(message, state)


@dp.message(Admin.add_photo, F.text == "-")
async def add_photo_skip(message: Message, state: FSMContext):
    await state.update_data(photo_id=None)
    await finish_add_event(message, state)


async def finish_add_event(message: Message, state: FSMContext):
    data = await state.get_data()
    event_id = f"evt_{uuid.uuid4().hex[:8]}"
    event = {
        "id": event_id,
        "title": data["title"],
        "date_text": data["date_text"],
        "start_dt": data["start_dt"],
        "place": data["place"],
        "description": data["description"],
        "price_single": data["price_single"],
        "price_pair": data["price_pair"],
        "photo_id": data.get("photo_id"),
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "active": True,
        "archived": False,
        "archived_at": None,
    }
    events.append(event)
    save_event(event_id)
    await state.clear()
    await message.answer(
        f"✅ Событие «{data['title']}» добавлено!",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🛠 Админ-панель", callback_data="admin:back")],
            [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
        ]),
    )


# ================== СПИСОК СОБЫТИЙ (АКТИВНЫЕ) ==================
@dp.callback_query(F.data == "admin:list")
async def admin_list(call: CallbackQuery):
    active = get_active_events()
    archived = get_archived_events()

    text = (
        f"📋 <b>Список событий</b>\n\n"
        f"🟢 Активных: <b>{len(active)}</b>\n"
        f"📦 В архиве: <b>{len(archived)}</b>\n\n"
    )

    if not active:
        text += "📭 Активных событий нет."
    else:
        text += "<b>Активные события:</b>\n\n"
        for e in active:
            ev_tickets = [t for t in tickets.values() if str(t.get("event_id")) == str(e["id"])]
            total = len(ev_tickets)
            active_cnt = sum(1 for t in ev_tickets if not t.get("used") and not t.get("auto_used"))
            used_cnt = sum(1 for t in ev_tickets if t.get("used") and not t.get("auto_used"))
            text += (
                f"🟢 <b>{e['title']}</b>\n"
                f"   📅 {e['date_text']}\n"
                f"   🎫 Билетов: {total} (🟢{active_cnt} 🔴{used_cnt})\n\n"
            )

    buttons = []
    for e in active:
        cnt = sum(1 for t in tickets.values() if str(t.get("event_id")) == str(e["id"]))
        buttons.append([InlineKeyboardButton(
            text=f"🎟 {e['title'][:35]} ({cnt})",
            callback_data=f"admin:event_tickets:{e['id']}:all:0"
        )])
    buttons.append([InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:list")])
    buttons.append([InlineKeyboardButton(text="📦 Архив событий", callback_data="admin:archive")])
    buttons.append([InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")])

    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")


# ================== УДАЛЕНИЕ СОБЫТИЙ ==================
@dp.callback_query(F.data == "admin:delete")
async def admin_delete(call: CallbackQuery, state: FSMContext):
    active = get_active_events()
    archived = get_archived_events()

    text = (
        f"🗑 <b>Удаление событий</b>\n\n"
        f"🟢 Активных: <b>{len(active)}</b>\n"
        f"📦 В архиве: <b>{len(archived)}</b>\n\n"
        f"Из какой категории удалить?"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"🟢 Из активных ({len(active)})", callback_data="admin:del_active")],
        [InlineKeyboardButton(text=f"📦 Из архива ({len(archived)})", callback_data="admin:del_archive")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])

    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:del_active")
async def admin_delete_active_list(call: CallbackQuery, state: FSMContext):
    active = get_active_events()

    if not active:
        await call.answer("Нет активных событий", show_alert=True)
        return

    buttons = []
    for e in active:
        buttons.append([InlineKeyboardButton(
            text=f"🗑 {e['title']} — {e['date_text']}",
            callback_data=f"admin:del_ev:{e['id']}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:delete")])

    await call.message.edit_text(
        "🗑 <b>Активные события</b> — выберите для удаления:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


@dp.callback_query(F.data == "admin:del_archive")
async def admin_delete_archive_list(call: CallbackQuery, state: FSMContext):
    archived = get_archived_events()

    if not archived:
        await call.answer("Архив пуст", show_alert=True)
        return

    archived_sorted = sorted(archived, key=lambda e: e.get("start_dt", ""), reverse=True)
    buttons = []
    for e in archived_sorted:
        buttons.append([InlineKeyboardButton(
            text=f"🗑 📦 {e['title']} — {e['date_text']}",
            callback_data=f"admin:del_ev:{e['id']}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:delete")])

    await call.message.edit_text(
        "🗑 <b>Архив</b> — выберите для удаления:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("admin:del_ev:"))
async def admin_delete_confirm(call: CallbackQuery, state: FSMContext):
    eid = call.data.split(":", 2)[2]
    events[:] = [e for e in events if e["id"] != eid]
    supabase.table("events").delete().eq("id", eid).execute()
    await call.message.edit_text("✅ Удалено.")
    await show_admin_menu(call.message, state)


@dp.callback_query(F.data.startswith("admin:del_arch:"))
async def admin_delete_arch_confirm(call: CallbackQuery, state: FSMContext):
    eid = call.data.split(":", 2)[2]
    events[:] = [e for e in events if e["id"] != eid]
    supabase.table("events").delete().eq("id", eid).execute()
    await call.message.edit_text("✅ Удалено из архива.")
    await show_admin_menu(call.message, state)


# ========== КОНЕЦ ЧАСТИ 9 ==========

# ================== БИЛЕТЫ ПО СОБЫТИЮ (ГЛАВНЫЙ ЭКРАН) ==================

TICKETS_BY_EVENT_PAGE_SIZE = 15


def _build_event_tickets_text(event: dict, filtered: list, filter_mode: str, page: int):
    """Собирает текст и кнопки для экрана билетов события."""
    total_pages = max(1, (len(filtered) + TICKETS_BY_EVENT_PAGE_SIZE - 1) // TICKETS_BY_EVENT_PAGE_SIZE)
    if page < 0:
        page = 0
    if page >= total_pages:
        page = total_pages - 1

    start = page * TICKETS_BY_EVENT_PAGE_SIZE
    end = start + TICKETS_BY_EVENT_PAGE_SIZE
    page_items = filtered[start:end]

    # Общая статистика по событию
    all_event_tickets = [
        (c, t) for c, t in tickets.items()
        if str(t.get("event_id")) == str(event["id"])
    ]
    total = len(all_event_tickets)
    active = sum(1 for _, t in all_event_tickets if not t.get("used") and not t.get("auto_used"))
    used = sum(1 for _, t in all_event_tickets if t.get("used") and not t.get("auto_used"))
    burned = sum(1 for _, t in all_event_tickets if t.get("auto_used"))
    revenue = sum(
        int(o.get("price", 0))
        for o in orders
        if str(o.get("event_id")) == str(event["id"])
    )

    filter_labels = {
        "all": "все",
        "active": "🟢 активные",
        "used": "🔴 погашенные",
        "burned": "⚪ сгоревшие",
    }
    filter_label = filter_labels.get(filter_mode, "все")

    text = (
        f"🎟 <b>{event['title']}</b>\n"
        f"📅 {event['date_text']}\n"
        f"📍 {event.get('place', '—')}\n\n"
        f"Всего: <b>{total}</b>\n"
        f"🟢 Активных: <b>{active}</b>\n"
        f"🔴 Погашено: <b>{used}</b>\n"
        f"⚪ Сгорело: <b>{burned}</b>\n"
        f"💰 Выручка: <b>{revenue} ₽</b>\n\n"
        f"Фильтр: <b>{filter_label}</b> · Найдено: <b>{len(filtered)}</b>\n"
        f"Страница {page + 1}/{total_pages}\n\n"
    )

    if not page_items:
        text += "📭 Билетов по фильтру нет.\n"
    else:
        for code, t in page_items:
            # Статус
            if t.get("auto_used"):
                status = "⚪ сгорел"
            elif t.get("used"):
                status = "🔴 погашен"
            else:
                status = "🟢 активен"

            # Имя
            holder = t.get("holder") or "—"

            # Телефон и username
            owner = users.get(t.get("user_id"), {})
            phone = owner.get("phone") or "—"
            username = owner.get("username") or ""

            # Дата покупки
            issued_at = t.get("issued_at") or "—"

            # Кто погасил
            used_by = t.get("used_by") or ""

            # Цена
                        # Цена
            price = 0
            paid_by_name = ""
            for o in orders:
                if (str(o.get("user_id")) == str(t.get("user_id"))
                        and str(o.get("event_id")) == str(t.get("event_id"))):
                    try:
                        price = int(o.get("price", 0) or 0)
                    except (TypeError, ValueError):
                        price = 0
                    break

            # Для парного билета без своего заказа — берём точного партнёра по partner_code
            if price == 0 and t.get("type") == "pair":
                partner_code = t.get("partner_code") or ""
                if partner_code:
                    partner_ticket = tickets.get(partner_code)
                    if partner_ticket:
                        partner_uid = partner_ticket.get("user_id")
                        partner_user = users.get(int(partner_uid or 0), {})
                        partner_uname = partner_user.get("username") or ""
                        if partner_uname:
                            paid_by_name = f" (оплачен @{partner_uname})"

            # Тип
            type_label = "Парный" if t.get("type") == "pair" else "Обычный"

            text += (
                f"{status} · <code>{code}</code>\n"
                f"   👤 {holder}"
            )
            if username:
                text += f" (@{username})"
            text += "\n"
            text += f"   📱 {phone}\n"
            text += f"   🎟 {type_label} · 💰 {price} ₽{paid_by_name}\n"
            text += f"   🛒 Куплен: {issued_at}\n"
            if used_by:
                text += f"   👮 Погасил: {used_by}\n"
            text += "\n"

    # Кнопки
    buttons = []

    # Фильтры
    buttons.append([
        InlineKeyboardButton(
            text="🟢 Активные" if filter_mode != "active" else "· 🟢 Активные ·",
            callback_data=f"admin:event_tickets:{event['id']}:active:0"
        ),
        InlineKeyboardButton(
            text="🔴 Погашенные" if filter_mode != "used" else "· 🔴 Погашенные ·",
            callback_data=f"admin:event_tickets:{event['id']}:used:0"
        ),
    ])
    buttons.append([
        InlineKeyboardButton(
            text="⚪ Сгоревшие" if filter_mode != "burned" else "· ⚪ Сгоревшие ·",
            callback_data=f"admin:event_tickets:{event['id']}:burned:0"
        ),
        InlineKeyboardButton(
            text="📋 Все" if filter_mode != "all" else "· 📋 Все ·",
            callback_data=f"admin:event_tickets:{event['id']}:all:0"
        ),
    ])

    # Пагинация
    if total_pages > 1:
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton(
                text="⬅️",
                callback_data=f"admin:event_tickets:{event['id']}:{filter_mode}:{page-1}"
            ))
        nav.append(InlineKeyboardButton(
            text=f"{page+1}/{total_pages}",
            callback_data="admin:noop"
        ))
        if page < total_pages - 1:
            nav.append(InlineKeyboardButton(
                text="➡️",
                callback_data=f"admin:event_tickets:{event['id']}:{filter_mode}:{page+1}"
            ))
        buttons.append(nav)

    # Действия
    buttons.append([
        InlineKeyboardButton(
            text="📥 XLSX",
            callback_data=f"admin:export_event_xlsx:{event['id']}:{filter_mode}"
        ),
        InlineKeyboardButton(
            text="📄 PDF",
            callback_data=f"admin:export_event_pdf:{event['id']}:{filter_mode}"
        ),
    ])

    buttons.append([InlineKeyboardButton(
        text="📋 Список на вход",
        callback_data=f"admin:checker_export_menu:ev:{event['id']}"
    )])

    buttons.append([InlineKeyboardButton(
        text="⬅️ К событиям", callback_data="admin:list"
    )])

    buttons.append([InlineKeyboardButton(
        text="🏠 На главную", callback_data="main_menu"
    )])

    return text, buttons, total_pages


@dp.callback_query(F.data.startswith("admin:event_tickets:"))
async def admin_event_tickets(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    try:
        event_id = parts[2]
    except IndexError:
        await call.answer("Ошибка данных", show_alert=True)
        return

    filter_mode = parts[3] if len(parts) > 3 else "all"
    page = int(parts[4]) if len(parts) > 4 and parts[4].isdigit() else 0

    event = next((e for e in events if str(e["id"]) == str(event_id)), None)
    if not event:
        await call.answer("Событие не найдено", show_alert=True)
        return

    # Все билеты события
    all_event_tickets = [
        (c, t) for c, t in tickets.items()
        if str(t.get("event_id")) == str(event_id)
    ]

    # Фильтр
    if filter_mode == "active":
        filtered = [(c, t) for c, t in all_event_tickets
                    if not t.get("used") and not t.get("auto_used")]
    elif filter_mode == "used":
        filtered = [(c, t) for c, t in all_event_tickets
                    if t.get("used") and not t.get("auto_used")]
    elif filter_mode == "burned":
        filtered = [(c, t) for c, t in all_event_tickets
                    if t.get("auto_used")]
    else:
        filtered = all_event_tickets

    # ✅ Сортировка: сначала активные, потом погашенные, потом сгоревшие.
    # Защита от None — используем `or ""`, иначе Python падает на
    # сравнении None < str (это был баг с TypeError в старом коде).
    def sort_key(kv):
        _, t = kv
        used_at = t.get("used_at") or ""
        issued_at = t.get("issued_at") or ""
        if t.get("auto_used"):
            return (2, used_at)
        if t.get("used"):
            return (1, used_at)
        return (0, issued_at)

    filtered.sort(key=sort_key, reverse=False)

    text, buttons, _ = _build_event_tickets_text(event, filtered, filter_mode, page)

    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    try:
        await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")
    except Exception:
        await call.message.answer(text[:4000], reply_markup=kb, parse_mode="HTML")
    await call.answer()


@dp.callback_query(F.data == "admin:tickets_by_event")
async def admin_tickets_by_event(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    if not events:
        await call.message.edit_text(
            "📭 Событий пока нет.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")]
            ]),
        )
        return

    # Активные + архивные
    buttons = []
    active = get_active_events()
    archived = get_archived_events()

    text = "🎯 <b>Билеты по событиям</b>\n\n"
    text += "Выберите событие:\n\n"

    if active:
        for e in active:
            cnt = sum(1 for t in tickets.values() if str(t.get("event_id")) == str(e["id"]))
            buttons.append([InlineKeyboardButton(
                text=f"🟢 {e['title'][:30]} — {e['date_text'][:12]} (🎟{cnt})",
                callback_data=f"admin:event_tickets:{e['id']}:all:0"
            )])

    if archived:
        archived_sorted = sorted(archived, key=lambda e: e.get("start_dt", ""), reverse=True)
        for e in archived_sorted:
            cnt = sum(1 for t in tickets.values() if str(t.get("event_id")) == str(e["id"]))
            buttons.append([InlineKeyboardButton(
                text=f"📦 {e['title'][:30]} — {e['date_text'][:12]} (🎟{cnt})",
                callback_data=f"admin:event_tickets:{e['id']}:all:0"
            )])

    buttons.append([InlineKeyboardButton(text="⬅️ К билетам", callback_data="admin:tickets:0")])
    buttons.append([InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")])

    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")
    await call.answer()


# ========== КОНЕЦ ЧАСТИ 10 ==========

# ================== СТАТИСТИКА (АДМИН) ==================
@dp.callback_query(F.data == "admin:stats")
async def admin_stats(call: CallbackQuery):
    total_revenue = sum(o["price"] for o in orders)
    paid_orders = len(orders)
    tickets_active = sum(1 for t in tickets.values() if not t.get("used") and not t.get("auto_used"))
    tickets_used = sum(1 for t in tickets.values() if t.get("used") and not t.get("auto_used"))
    tickets_auto_used = sum(1 for t in tickets.values() if t.get("auto_used"))

    by_event = {}
    for o in orders:
        title = o.get("event_title", "—")
        by_event.setdefault(title, {"count": 0, "sum": 0})
        by_event[title]["count"] += 1
        by_event[title]["sum"] += o["price"]

    active_events = get_active_events()
    archived_events = get_archived_events()

    text = (
        f"📊 <b>Статистика</b>\n\n"
        f"👥 Пользователей: <b>{len(users)}</b>\n"
        f"🎉 Событий активно: <b>{len(active_events)}</b>\n"
        f"📦 В архиве: <b>{len(archived_events)}</b>\n"
        f"🛒 Заказов: <b>{paid_orders}</b>\n"
        f"💰 Выручка: <b>{total_revenue} ₽</b>\n\n"
        f"🟢 Билетов активно: <b>{tickets_active}</b>\n"
        f"🔴 Погашено: <b>{tickets_used}</b>\n"
        f"⚪ Сгорело: <b>{tickets_auto_used}</b>\n"
        f"🚫 В ЧС: <b>{len(blacklist)}</b>\n"
        f"📊 Попыток рег.: <b>{len(reg_attempts)}</b>\n"
        f"🎫 Checkers: <b>{len(checkers)}</b>\n"
    )

    if by_event:
        text += "\n<b>Продажи по событиям:</b>\n"
        for title, info in by_event.items():
            text += f"• {title}: {info['count']} шт, {info['sum']} ₽\n"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:stats")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


# ================== ПОЛЬЗОВАТЕЛИ (АДМИН) ==================
PAGE_SIZE = 10


@dp.callback_query(F.data.startswith("admin:users:"))
async def admin_users(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    page = int(call.data.split(":")[2])
    all_ids = list(users.keys())
    total = len(all_ids)
    total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)

    if page < 0:
        page = 0
    if page >= total_pages:
        page = total_pages - 1

    start = page * PAGE_SIZE
    end = start + PAGE_SIZE
    page_ids = all_ids[start:end]

    text = f"👥 <b>Пользователи</b> ({total})\n"
    text += f"Страница {page + 1}/{total_pages}\n\n"

    for uid in page_ids:
        u = users[uid]
        text += (
            f"👤 <b>{u.get('name','')} {u.get('surname','')}</b>\n"
            f"🆔 <code>{uid}</code>\n"
            f"💬 @{u.get('username','—') or '—'}\n"
            f"📱 {u.get('phone','—')}\n\n"
        )

    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton(text="⬅️ Назад", callback_data=f"admin:users:{page-1}"))
    nav.append(InlineKeyboardButton(text=f"{page+1}/{total_pages}", callback_data="admin:noop"))
    if page < total_pages - 1:
        nav.append(InlineKeyboardButton(text="Вперёд ➡️", callback_data=f"admin:users:{page+1}"))

    kb = InlineKeyboardMarkup(inline_keyboard=[
        nav,
        [InlineKeyboardButton(text="📥 Скачать CSV", callback_data="admin:users_csv")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:noop")
async def admin_noop(call: CallbackQuery):
    await call.answer()


@dp.callback_query(F.data == "admin:users_csv")
async def admin_users_csv(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    rows = [[
        uid, u.get("username", ""), u.get("email", ""),
        u.get("name", ""), u.get("surname", ""), u.get("gender", ""),
        u.get("phone", ""), u.get("registered_at", ""),
    ] for uid, u in users.items()]

    csv_bytes = make_csv(
        ["telegram_id", "username", "email", "name", "surname",
         "gender", "phone", "registered_at"],
        rows
    )
    await call.message.answer_document(
        BufferedInputFile(csv_bytes, filename="users.csv"),
        caption=f"👥 Пользователи ({len(rows)})",
    )
    await call.answer()


# ================== РАССЫЛКА ==================
@dp.callback_query(F.data == "admin:broadcast")
async def admin_broadcast_start(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    await call.message.edit_text(
        f"📢 <b>Рассылка</b>\n\n"
        f"Получателей: <b>{len(users)}</b>\n\n"
        f"Отправьте текст сообщения (можно с фото):\n\n"
        f"Для отмены — /admin",
        parse_mode="HTML",
    )
    await state.set_state(Admin.broadcast)


@dp.message(Admin.broadcast)
async def admin_broadcast_send(message: Message, state: FSMContext):
    if message.text == "/admin":
        await state.clear()
        await show_admin_menu(message, state)
        return

    await state.clear()
    await message.answer("📢 Начинаю рассылку...")

    sent = 0
    failed = 0
    for uid in list(users.keys()):
        try:
            if message.photo:
                await bot.send_photo(
                    uid,
                    message.photo[-1].file_id,
                    caption=message.caption or "",
                )
            else:
                await bot.send_message(uid, message.text or "")
            sent += 1
            await asyncio.sleep(0.05)
        except Exception:
            failed += 1

    await message.answer(
        f"✅ Рассылка завершена\n\n"
        f"📤 Отправлено: <b>{sent}</b>\n"
        f"❌ Ошибок: <b>{failed}</b>",
        parse_mode="HTML",
    )
    await show_admin_menu(message, state)


# ================== НАЙТИ БИЛЕТ (АДМИН) ==================
@dp.callback_query(F.data == "admin:find_ticket")
async def admin_find_ticket_start(call: CallbackQuery, state: FSMContext):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    await call.message.edit_text(
        "🔍 <b>Поиск билета</b>\n\n"
        "Введите <b>код</b>, <b>имя</b>, <b>телефон</b> или <b>email</b>:",
        parse_mode="HTML",
    )
    await state.set_state(Admin.find_ticket)


@dp.message(Admin.find_ticket)
async def admin_find_ticket_query(message: Message, state: FSMContext):
    query = message.text.strip().lower()
    if len(query) < 2:
        await message.answer("❌ Слишком короткий запрос.")
        return

    found = []
    for code, t in tickets.items():
        holder = (t.get("holder") or "").lower()
        owner = users.get(t.get("user_id"), {})
        owner_name = f"{owner.get('name','')} {owner.get('surname','')}".strip().lower()
        owner_phone = (owner.get("phone") or "").replace(" ", "").replace("-", "").replace("+", "")
        owner_email = (owner.get("email") or "").lower()

        if (query in code.lower()
                or query in holder
                or query in owner_name
                or query in owner_phone
                or query in owner_email):
            found.append((code, t))

    if not found:
        await message.answer(
            "❌ Ничего не найдено.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
            ]),
        )
        await state.clear()
        return

    if len(found) == 1:
        code, t = found[0]
        await show_super_admin_ticket_card(message, code, t)
        await state.clear()
        return

    buttons = []
    for code, t in found[:30]:
        holder = t.get("holder", "—")
        status = get_ticket_status_icon(t)
        buttons.append([InlineKeyboardButton(
            text=f"{status} {code} — {holder[:20]}",
            callback_data=f"admin:ticket:show:{code}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")])

    await message.answer(
        f"🔍 Найдено: <b>{len(found)}</b>\n\nВыберите:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("admin:ticket:show:"))
async def admin_ticket_show(call: CallbackQuery):
    code = call.data.split(":", 3)[3]
    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Не найден", show_alert=True)
        return
    await call.answer()
    await show_super_admin_ticket_card(call.message, code, ticket)


# ================== ВСЕ БИЛЕТЫ (АДМИН) ==================
TICKETS_PAGE_SIZE = 15


@dp.callback_query(F.data.startswith("admin:tickets"))
async def admin_tickets(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    page = int(parts[2]) if len(parts) > 2 and parts[2].isdigit() else 0

    all_items = [(c, t) for c, t in tickets.items() if not t.get("auto_used")]
    # ✅ Защита от None в issued_at
    all_items.sort(key=lambda kv: kv[1].get("issued_at") or "", reverse=True)

    total = len(all_items)
    total_pages = max(1, (total + TICKETS_PAGE_SIZE - 1) // TICKETS_PAGE_SIZE)

    if page < 0:
        page = 0
    if page >= total_pages:
        page = total_pages - 1

    start = page * TICKETS_PAGE_SIZE
    end = start + TICKETS_PAGE_SIZE
    page_items = all_items[start:end]

    active = sum(1 for _, t in all_items if not t.get("used"))
    used = total - active

    text = (
        f"🎟 <b>Все билеты</b>\n\n"
        f"Всего: <b>{total}</b>\n"
        f"🟢 Активных: <b>{active}</b>\n"
        f"🔴 Погашено: <b>{used}</b>\n\n"
        f"Страница {page + 1}/{total_pages}\n\n"
    )

    for code, t in page_items:
        status = "🔴" if t.get("used") else "🟢"
        holder = t.get("holder", "—")
        title = t.get("event_title", "—")
        date_text = t.get("event_date", "—")
        text += (
            f"{status} <code>{code}</code>\n"
            f"   👤 {holder}\n"
            f"   🎉 {title} — {date_text}\n\n"
        )

    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton(text="⬅️", callback_data=f"admin:tickets:{page-1}"))
    nav.append(InlineKeyboardButton(text=f"{page+1}/{total_pages}", callback_data="admin:noop"))
    if page < total_pages - 1:
        nav.append(InlineKeyboardButton(text="➡️", callback_data=f"admin:tickets:{page+1}"))

    kb = InlineKeyboardMarkup(inline_keyboard=[
        nav,
        [InlineKeyboardButton(text="🎯 По событиям", callback_data="admin:tickets_by_event")],
        [InlineKeyboardButton(text="🔄 Обновить", callback_data=f"admin:tickets:{page}")],
        [InlineKeyboardButton(text="📥 Скачать CSV", callback_data="admin:csv_all_tickets")],
        [InlineKeyboardButton(text="📤 Выгрузка билетов", callback_data="admin:export")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:csv_all_tickets")
async def admin_tickets_csv(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

        rows = []
    for code, t in tickets.items():
        owner = users.get(t.get("user_id"), {})
        status = "сгорел" if t.get("auto_used") else ("использован" if t.get("used") else "активен")

        # Цена
        price = 0
        paid_by = ""
        for o in orders:
            if (str(o.get("user_id")) == str(t.get("user_id"))
                    and str(o.get("event_id")) == str(t.get("event_id"))):
                try:
                    price = int(o.get("price", 0) or 0)
                except (TypeError, ValueError):
                    price = 0
                break

        if price == 0 and t.get("type") == "pair":
            partner_code = t.get("partner_code") or ""
            if partner_code:
                partner_ticket = tickets.get(partner_code)
                if partner_ticket:
                    partner_uid = partner_ticket.get("user_id")
                    partner_user = users.get(int(partner_uid or 0), {})
                    partner_uname = partner_user.get("username") or ""
                    if partner_uname:
                        paid_by = f"@{partner_uname}"

        rows.append([
            code,
            t.get("user_id", ""),
            t.get("holder", ""),
            owner.get("phone", ""),
            owner.get("username", ""),
            t.get("type", ""),
            status,
            t.get("event_title", ""),
            t.get("event_date", ""),
            t.get("issued_at", ""),
            t.get("used_at", ""),
            t.get("used_by", ""),
            price,
            paid_by,
        ])

    csv_bytes = make_csv(
        ["ticket_code", "holder_id", "holder_name", "phone", "username",
         "type", "status", "event_title", "event_date", "issued_at",
         "used_at", "used_by"],
        rows
    )
    await call.message.answer_document(
        BufferedInputFile(csv_bytes, filename="tickets.csv"),
        caption=f"🎫 Все билеты ({len(rows)})",
    )
    await call.answer()


@dp.callback_query(F.data == "admin:transfers")
async def admin_transfers(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    if not ticket_transfers:
        await call.message.edit_text(
            "📭 История передач пуста.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")]
            ]),
        )
        return

    text = f"📋 <b>История передач</b> ({len(ticket_transfers)})\n\n"

    for tr in ticket_transfers[:20]:
        from_name = tr.get("from_name") or "—"
        from_uname = tr.get("from_username") or ""
        to_name = tr.get("to_name") or "—"
        to_uname = tr.get("to_username") or ""

        text += (
            f"🎫 <code>{tr.get('ticket_code', '—')}</code>\n"
            f"   👤 От: {from_name}"
        )
        if from_uname:
            text += f" (@{from_uname})"
        text += f"\n   👤 Кому: {to_name}"
        if to_uname:
            text += f" (@{to_uname})"

        tr_price = 0
        tr_type = ""
        tr_ticket = tickets.get(tr.get("ticket_code", ""))
        if tr_ticket:
            tr_type = "Парный" if tr_ticket.get("type") == "pair" else "Обычный"
            for o in orders:
                if (str(o.get("user_id")) == str(tr.get("from_user_id"))
                        and str(o.get("event_id")) == str(tr.get("event_id"))):
                    try:
                        tr_price = int(o.get("price", 0) or 0)
                    except (TypeError, ValueError):
                        tr_price = 0
                    break

        if tr_type:
            text += f"\n   🎫 {tr_type}"
            if tr_price > 0:
                text += f" · 💰 {tr_price} ₽"

        text += (
            f"\n   🎉 {tr.get('event_title', '—')}\n"
            f"   🕐 {tr.get('transferred_at', '—')}\n\n"
        )

    if len(ticket_transfers) > 20:
        text += f"... и ещё {len(ticket_transfers) - 20}\n"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:transfers")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])

    try:
        await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")
    except Exception:
        await call.message.answer(text[:4000], reply_markup=kb, parse_mode="HTML")
    await call.answer()


# ========== КОНЕЦ ЧАСТИ 11 ==========

# ================== АРХИВ БИЛЕТОВ (АДМИН) ==================
@dp.callback_query(F.data == "admin:archived_tickets")
async def admin_archived_tickets(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    archived = []
    for code, t in tickets.items():
        if not t.get("auto_used"):
            continue
        if is_ticket_archived(t):
            archived.append((code, t))

    # ✅ Защита от None
    archived.sort(key=lambda kv: kv[1].get("used_at") or "", reverse=True)

    if not archived:
        await call.message.edit_text(
            "📭 Архив билетов пуст.\n\n"
            "<i>Билеты попадают сюда через 30 дней после авто-погашения.</i>",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")]
            ]),
            parse_mode="HTML",
        )
        return

    text = f"📦 <b>Архив билетов</b> ({len(archived)})\n\n"
    for code, t in archived[:20]:
        text += (
            f"⚪ <code>{code}</code>\n"
            f"   👤 {t.get('holder', '—')}\n"
            f"   🎉 {t.get('event_title', '—')}\n"
            f"   📅 {t.get('event_date', '—')}\n\n"
        )
    if len(archived) > 20:
        text += f"... и ещё {len(archived) - 20}"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📥 Скачать CSV", callback_data="admin:archived_tickets_csv")],
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:archived_tickets")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text[:4000], reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:archived_tickets_csv")
@dp.callback_query(F.data == "admin:archived_tickets_csv")
async def admin_archived_tickets_csv(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    rows = []
    for code, t in tickets.items():
        if not t.get("auto_used"):
            continue
        if not is_ticket_archived(t):
            continue
        owner = users.get(t.get("user_id"), {})

        # Цена
        price = 0
        paid_by = ""
        for o in orders:
            if (str(o.get("user_id")) == str(t.get("user_id"))
                    and str(o.get("event_id")) == str(t.get("event_id"))):
                try:
                    price = int(o.get("price", 0) or 0)
                except (TypeError, ValueError):
                    price = 0
                break

        if price == 0 and t.get("type") == "pair":
            partner_code = t.get("partner_code") or ""
            if partner_code:
                partner_ticket = tickets.get(partner_code)
                if partner_ticket:
                    partner_uid = partner_ticket.get("user_id")
                    partner_user = users.get(int(partner_uid or 0), {})
                    partner_uname = partner_user.get("username") or ""
                    if partner_uname:
                        paid_by = f"@{partner_uname}"

        rows.append([
            code,
            t.get("user_id", ""),
            t.get("holder", ""),
            owner.get("phone", ""),
            owner.get("username", ""),
            t.get("type", ""),
            "сгорел",
            t.get("event_title", ""),
            t.get("event_date", ""),
            t.get("issued_at", ""),
            t.get("used_at", ""),
            t.get("used_by", ""),
            price,
            paid_by,
        ])

    csv_bytes = make_csv(
        ["ticket_code", "holder_id", "holder_name", "phone", "username",
         "type", "status", "event_title", "event_date", "issued_at",
         "used_at", "used_by", "price", "paid_by"],
        rows
    )
    await call.message.answer_document(
        BufferedInputFile(csv_bytes, filename="archived_tickets.csv"),
        caption=f"📦 Архив билетов ({len(rows)})",
    )
    await call.answer()

# ================== ПРОСТАЯ ВЫГРУЗКА ДЛЯ CHECKER'ОВ ==================


@dp.callback_query(F.data.startswith("admin:checker_export_menu:"))
async def admin_checker_export_menu(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    # admin:checker_export_menu:all:  или  admin:checker_export_menu:ev:<event_id>
    scope = parts[2] if len(parts) > 2 else "all"
    event_id = parts[3] if len(parts) > 3 else ""

    # Считаем активные билеты (для этого события или для всех)
    count = 0
    for code, t in tickets.items():
        if t.get("used") or t.get("auto_used"):
            continue
        if event_id and str(t.get("event_id")) != str(event_id):
            continue
        count += 1

    title = "📋 <b>Список на вход</b>\n\n"
    if event_id:
        ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
        if ev:
            title += f"🎉 {ev['title']}\n📅 {ev['date_text']}\n"
    else:
        title += "🌐 Все события\n"
    title += f"\nАктивных билетов: <b>{count}</b>\n\nВыберите формат:"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text="📊 Excel (XLSX)",
            callback_data=f"admin:checker_export_xlsx:{scope}:{event_id}"
        )],
        [InlineKeyboardButton(
            text="📄 PDF",
            callback_data=f"admin:checker_export_pdf:{scope}:{event_id}"
        )],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:back")],
    ])

    try:
        await call.message.edit_text(title, reply_markup=kb, parse_mode="HTML")
    except Exception:
        await call.message.answer(title, reply_markup=kb, parse_mode="HTML")
    await call.answer()

def _build_checkers_rows(event_id: str = "") -> list:
    """Собирает строки для выгрузки checker'ов."""
    rows = []
    for code, t in tickets.items():
        if t.get("used") or t.get("auto_used"):
            continue
        if event_id and str(t.get("event_id")) != str(event_id):
            continue

        owner = users.get(t.get("user_id"), {})
        name = (owner.get("name") or "").strip()
        surname = (owner.get("surname") or "").strip()
        holder = f"{name} {surname}".strip() or t.get("holder") or "—"

        price = 0
        for o in orders:
            if (str(o.get("user_id")) == str(t.get("user_id"))
                    and str(o.get("event_id")) == str(t.get("event_id"))):
                try:
                    price = int(o.get("price", 0) or 0)
                except (TypeError, ValueError):
                    price = 0
                break

        rows.append([code, holder, t.get("event_date", ""), price, "активен"])
    return rows


def _checkers_headers_ru():
    return ["Код билета", "Гость", "Дата вечеринки", "Цена", "Статус"]


@dp.callback_query(F.data.startswith("admin:checker_export_xlsx:"))
async def admin_checker_export_xlsx(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    event_id = parts[3] if len(parts) > 3 else ""
    rows = _build_checkers_rows(event_id)

    if not rows:
        await call.answer("Нет активных билетов", show_alert=True)
        return

    xlsx_bytes = make_xlsx(_checkers_headers_ru(), rows)
    if event_id:
        ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
        safe_name = _safe_filename(ev["title"]) if ev else "event"
        filename = f"checkers_{safe_name}.xlsx"
    else:
        filename = "checkers_all.xlsx"

    await call.message.answer_document(
        BufferedInputFile(xlsx_bytes, filename=filename),
        caption=f"📋 Список на вход (XLSX): {len(rows)} билетов",
    )
    await call.answer()


@dp.callback_query(F.data.startswith("admin:checker_export_pdf:"))
async def admin_checker_export_pdf(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    event_id = parts[3] if len(parts) > 3 else ""
    rows = _build_checkers_rows(event_id)

    if not rows:
        await call.answer("Нет активных билетов", show_alert=True)
        return

    title = "Список на вход"
    if event_id:
        ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
        if ev:
            title = f"{ev['title']} — список на вход"

    pdf_bytes = make_pdf(_checkers_headers_ru(), rows, title=title)
    if event_id:
        ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
        safe_name = _safe_filename(ev["title"]) if ev else "event"
        filename = f"checkers_{safe_name}.pdf"
    else:
        filename = "checkers_all.pdf"

    await call.message.answer_document(
        BufferedInputFile(pdf_bytes, filename=filename),
        caption=f"📋 Список на вход (PDF): {len(rows)} билетов",
    )
    await call.answer()


# ================== ФУНКЦИИ ЭКСПОРТА ==================
def _export_tickets_data(tickets_filtered: list) -> list:
    rows = []
    for code, t in tickets_filtered:
        owner = users.get(t.get("user_id"), {})
        if t.get("auto_used"):
            status = "сгорел"
        elif t.get("used"):
            status = "использован"
        else:
            status = "активен"

        # Цена: ищем заказ по user_id + event_id
        price = 0
        paid_by = ""
        for o in orders:
            if (str(o.get("user_id")) == str(t.get("user_id"))
                    and str(o.get("event_id")) == str(t.get("event_id"))):
                try:
                    price = int(o.get("price", 0) or 0)
                except (TypeError, ValueError):
                    price = 0
                break

        # Парный билет без своего заказа: показываем 0 + кто оплатил
        if price == 0 and t.get("type") == "pair":
            partner_code = t.get("partner_code") or ""
            if partner_code:
                partner_ticket = tickets.get(partner_code)
                if partner_ticket:
                    partner_uid = partner_ticket.get("user_id")
                    partner_user = users.get(int(partner_uid or 0), {})
                    partner_uname = partner_user.get("username") or ""
                    if partner_uname:
                        paid_by = f"@{partner_uname}"

        rows.append([
            code,
            t.get("user_id", ""),
            t.get("holder", ""),
            owner.get("phone", ""),
            owner.get("username", ""),
            t.get("type", ""),
            status,
            t.get("event_title", ""),
            t.get("event_date", ""),
            t.get("issued_at", ""),
            t.get("used_at", ""),
            t.get("used_by", ""),
            price,
            paid_by,
        ])
    return rows


def _tickets_headers():
    return [
        "ticket_code", "holder_id", "holder_name", "phone", "username",
        "type", "status", "event_title", "event_date", "issued_at",
        "used_at", "used_by", "price", "paid_by",
    ]


def make_xlsx(headers: list, rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Билеты"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_num, row in enumerate(rows, 2):
        for col_num, value in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    for col_num in range(1, len(headers) + 1):
        col_letter = ws.cell(row=1, column=col_num).column_letter
        max_len = len(str(headers[col_num - 1]))
        for row_num in range(2, len(rows) + 2):
            val = ws.cell(row=row_num, column=col_num).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def make_pdf(headers: list, rows: list, title: str = "Список билетов") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=15, rightMargin=15, topMargin=15, bottomMargin=15,
    )

    font_paths = [
        os.path.join(os.path.dirname(__file__), "DejaVuSans.ttf"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ]
    base_font = "Helvetica"
    for path in font_paths:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("DejaVu", path))
                base_font = "DejaVu"
                break
            except Exception:
                continue

    styles = getSampleStyleSheet()
    title_style = styles["Heading1"]
    title_style.fontName = base_font

    elements = []
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 12))

    table_data = [headers] + rows

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), base_font),
        ("FONTNAME", (0, 1), (-1, -1), base_font),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def make_zip_qr(tickets_filtered: list) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for code, t in tickets_filtered:
            qr = qrcode.QRCode(version=1, box_size=10, border=2)
            qr.add_data(code)
            qr.make(fit=True)
            img = qr.make_image(fill_color="black", back_color="white")

            img_buf = io.BytesIO()
            img.save(img_buf, format="PNG")
            img_buf.seek(0)

            holder = (t.get("holder") or "guest").replace(" ", "_")[:20]
            filename = f"{code}_{holder}.png"
            zf.writestr(filename, img_buf.getvalue())
    buf.seek(0)
    return buf.getvalue()


def _apply_filter(mode: str) -> list:
    if mode == "all":
        return list(tickets.items())
    elif mode == "used":
        return [(c, t) for c, t in tickets.items() if t.get("used") and not t.get("auto_used")]
    elif mode == "active":
        return [(c, t) for c, t in tickets.items() if not t.get("used") and not t.get("auto_used")]
    elif mode == "pair":
        return [(c, t) for c, t in tickets.items() if t.get("type") == "pair"]
    elif mode == "single":
        return [(c, t) for c, t in tickets.items() if t.get("type") != "pair"]
    elif mode == "burned":
        return [(c, t) for c, t in tickets.items() if t.get("auto_used")]
    return list(tickets.items())


# ================== ЭКСПОРТ ПО СОБЫТИЮ (с фильтром) ==================

def _get_event_filtered(event_id: str, filter_mode: str) -> list:
    """Возвращает билеты события по фильтру."""
    all_event_tickets = [
        (c, t) for c, t in tickets.items()
        if str(t.get("event_id")) == str(event_id)
    ]

    if filter_mode == "active":
        return [(c, t) for c, t in all_event_tickets
                if not t.get("used") and not t.get("auto_used")]
    elif filter_mode == "used":
        return [(c, t) for c, t in all_event_tickets
                if t.get("used") and not t.get("auto_used")]
    elif filter_mode == "burned":
        return [(c, t) for c, t in all_event_tickets
                if t.get("auto_used")]
    else:
        return all_event_tickets


def _safe_filename(text: str, max_len: int = 30) -> str:
    """Безопасное имя файла из названия события."""
    safe = "".join(c for c in text if c.isalnum() or c in " -_")
    safe = safe.strip().replace(" ", "_")
    return safe[:max_len] if safe else "event"


@dp.callback_query(F.data.startswith("admin:export_event_xlsx:"))
async def admin_export_event_xlsx(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    try:
        event_id = parts[2]
    except IndexError:
        await call.answer("Ошибка данных", show_alert=True)
        return

    filter_mode = parts[3] if len(parts) > 3 else "all"

    event = next((e for e in events if str(e["id"]) == str(event_id)), None)
    if not event:
        await call.answer("Событие не найдено", show_alert=True)
        return

    filtered = _get_event_filtered(event_id, filter_mode)

    if not filtered:
        await call.answer("Нет билетов по фильтру", show_alert=True)
        return

    # ✅ Защита от None при сортировке
    def sort_key(kv):
        _, t = kv
        used_at = t.get("used_at") or ""
        issued_at = t.get("issued_at") or ""
        if t.get("auto_used"):
            return (2, used_at)
        if t.get("used"):
            return (1, used_at)
        return (0, issued_at)

    filtered.sort(key=sort_key)

    filter_labels = {
        "all": "все",
        "active": "активные",
        "used": "погашенные",
        "burned": "сгоревшие",
    }
    filter_label = filter_labels.get(filter_mode, "все")

    rows = _export_tickets_data(filtered)
    xlsx_bytes = make_xlsx(_tickets_headers(), rows)

    safe_name = _safe_filename(event['title'])
    filename = f"tickets_{safe_name}_{filter_mode}.xlsx"

    await call.message.answer_document(
        BufferedInputFile(xlsx_bytes, filename=filename),
        caption=f"📊 Excel: {len(filtered)} билетов · {event['title']} ({filter_label})",
    )
    await call.answer()


@dp.callback_query(F.data.startswith("admin:export_event_pdf:"))
async def admin_export_event_pdf(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    parts = call.data.split(":")
    try:
        event_id = parts[2]
    except IndexError:
        await call.answer("Ошибка данных", show_alert=True)
        return

    filter_mode = parts[3] if len(parts) > 3 else "all"

    event = next((e for e in events if str(e["id"]) == str(event_id)), None)
    if not event:
        await call.answer("Событие не найдено", show_alert=True)
        return

    filtered = _get_event_filtered(event_id, filter_mode)

    if not filtered:
        await call.answer("Нет билетов по фильтру", show_alert=True)
        return

    # ✅ Защита от None при сортировке
    def sort_key(kv):
        _, t = kv
        used_at = t.get("used_at") or ""
        issued_at = t.get("issued_at") or ""
        if t.get("auto_used"):
            return (2, used_at)
        if t.get("used"):
            return (1, used_at)
        return (0, issued_at)

    filtered.sort(key=sort_key)

    filter_labels = {
        "all": "все",
        "active": "активные",
        "used": "погашенные",
        "burned": "сгоревшие",
    }
    filter_label = filter_labels.get(filter_mode, "все")

    pdf_title = f"{event['title']} ({filter_label})"
    if event.get("date_text"):
        pdf_title += f" — {event['date_text']}"

    rows = _export_tickets_data(filtered)
    pdf_bytes = make_pdf(_tickets_headers(), rows, title=pdf_title)

    safe_name = _safe_filename(event['title'])
    filename = f"tickets_{safe_name}_{filter_mode}.pdf"

    await call.message.answer_document(
        BufferedInputFile(pdf_bytes, filename=filename),
        caption=f"📄 PDF: {len(filtered)} билетов · {event['title']} ({filter_label})",
    )
    await call.answer()


# ========== КОНЕЦ ЧАСТИ 12 ==========


# ================== МЕНЮ ЭКСПОРТА (общее) ==================
@dp.callback_query(F.data == "admin:export")
async def admin_export_menu(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    total = len(tickets)
    active = sum(1 for t in tickets.values() if not t.get("used") and not t.get("auto_used"))
    used = sum(1 for t in tickets.values() if t.get("used") and not t.get("auto_used"))
    burned = sum(1 for t in tickets.values() if t.get("auto_used"))

    text = (
        f"📤 <b>Выгрузка билетов</b>\n\n"
        f"Всего: <b>{total}</b>\n"
        f"🟢 Активных: <b>{active}</b>\n"
        f"🔴 Погашено: <b>{used}</b>\n"
        f"⚪ Сгорело: <b>{burned}</b>\n\n"
        f"Выберите формат:"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📊 Excel (XLSX)", callback_data="admin:export:xlsx:all")],
        [InlineKeyboardButton(text="📄 PDF", callback_data="admin:export:pdf:all")],
        [InlineKeyboardButton(text="📦 ZIP с QR-кодами", callback_data="admin:export:zip:all")],
        [InlineKeyboardButton(text="📋 CSV", callback_data="admin:export:csv:all")],
        [InlineKeyboardButton(text="📋 Список на вход", callback_data="admin:checker_export_menu:all:")],
        [InlineKeyboardButton(text="🎯 Фильтры", callback_data="admin:export:filter")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])

    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:export:filter")
async def admin_export_filter(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔴 Только погашенные", callback_data="admin:export:xlsx:used")],
        [InlineKeyboardButton(text="🟢 Только активные", callback_data="admin:export:xlsx:active")],
        [InlineKeyboardButton(text="⚪ Только сгоревшие", callback_data="admin:export:xlsx:burned")],
        [InlineKeyboardButton(text="🎟 Только парные", callback_data="admin:export:xlsx:pair")],
        [InlineKeyboardButton(text="🎫 Только обычные", callback_data="admin:export:xlsx:single")],
        [InlineKeyboardButton(text="🎉 За конкретное событие", callback_data="admin:export:byevent")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:export")],
    ])

    await call.message.edit_text(
        "🎯 <b>Фильтры билетов</b>\n\n"
        "Выберите, какие билеты включить в выгрузку:",
        reply_markup=kb, parse_mode="HTML",
    )


@dp.callback_query(F.data == "admin:export:byevent")
async def admin_export_byevent(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    if not events:
        await call.answer("Нет событий", show_alert=True)
        return

    buttons = []
    for e in events:
        cnt = sum(1 for t in tickets.values() if str(t.get("event_id")) == str(e["id"]))
        buttons.append([InlineKeyboardButton(
            text=f"🎉 {e['title'][:30]} ({cnt})",
            callback_data=f"admin:export:event:{e['id']}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:export:filter")])

    await call.message.edit_text(
        "🎉 <b>Выберите событие</b>:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("admin:export:event:"))
async def admin_export_event(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    event_id = call.data.split(":", 3)[3]
    filtered = [(c, t) for c, t in tickets.items() if str(t.get("event_id")) == str(event_id)]

    if not filtered:
        await call.answer("Нет билетов на это событие", show_alert=True)
        return

    ev = next((e for e in events if str(e["id"]) == str(event_id)), None)
    title = ev["title"] if ev else "event"

    rows = _export_tickets_data(filtered)
    xlsx_bytes = make_xlsx(_tickets_headers(), rows)

    await call.message.answer_document(
        BufferedInputFile(xlsx_bytes, filename=f"tickets_{event_id}.xlsx"),
        caption=f"📊 Билеты на «{title}» ({len(filtered)} шт)",
    )
    await call.answer()


@dp.callback_query(F.data.startswith("admin:export:xlsx:"))
async def admin_export_xlsx(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    mode = call.data.split(":", 3)[3]
    filtered = _apply_filter(mode)

    if not filtered:
        await call.answer("Нет билетов по фильтру", show_alert=True)
        return

    rows = _export_tickets_data(filtered)
    xlsx_bytes = make_xlsx(_tickets_headers(), rows)

    await call.message.answer_document(
        BufferedInputFile(xlsx_bytes, filename=f"tickets_{mode}.xlsx"),
        caption=f"📊 Excel: {len(filtered)} билетов ({mode})",
    )
    await call.answer()


@dp.callback_query(F.data.startswith("admin:export:pdf:"))
async def admin_export_pdf(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    mode = call.data.split(":", 3)[3]
    filtered = _apply_filter(mode)

    if not filtered:
        await call.answer("Нет билетов по фильтру", show_alert=True)
        return

    rows = _export_tickets_data(filtered)
    pdf_bytes = make_pdf(_tickets_headers(), rows, title=f"Билеты ({mode}, {len(filtered)} шт)")

    await call.message.answer_document(
        BufferedInputFile(pdf_bytes, filename=f"tickets_{mode}.pdf"),
        caption=f"📄 PDF: {len(filtered)} билетов ({mode})",
    )
    await call.answer()


@dp.callback_query(F.data.startswith("admin:export:zip:"))
async def admin_export_zip(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    mode = call.data.split(":", 3)[3]
    filtered = _apply_filter(mode)

    if not filtered:
        await call.answer("Нет билетов по фильтру", show_alert=True)
        return

    if len(filtered) > 200:
        await call.answer(f"Слишком много билетов ({len(filtered)}). Максимум 200.", show_alert=True)
        return

    await call.answer("📦 Генерирую ZIP...")
    zip_bytes = make_zip_qr(filtered)

    await call.message.answer_document(
        BufferedInputFile(zip_bytes, filename=f"qr_codes_{mode}.zip"),
        caption=f"📦 ZIP с QR: {len(filtered)} билетов ({mode})",
    )


@dp.callback_query(F.data.startswith("admin:export:csv:"))
async def admin_export_csv_filtered(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    mode = call.data.split(":", 3)[3]
    filtered = _apply_filter(mode)

    if not filtered:
        await call.answer("Нет билетов по фильтру", show_alert=True)
        return

    rows = _export_tickets_data(filtered)
    csv_bytes = make_csv(_tickets_headers(), rows)

    await call.message.answer_document(
        BufferedInputFile(csv_bytes, filename=f"tickets_{mode}.csv"),
        caption=f"📋 CSV: {len(filtered)} билетов ({mode})",
    )
    await call.answer()




# ================== ЧЁРНЫЙ СПИСОК (АДМИН) ==================
@dp.callback_query(F.data == "admin:blacklist")
async def admin_blacklist_menu(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    text = get_blacklist_text()[:4000]

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Добавить username", callback_data="admin:bl:add_username")],
        [InlineKeyboardButton(text="📱 Добавить телефон", callback_data="admin:bl:add_phone")],
        [InlineKeyboardButton(text="🗑 Удалить", callback_data="admin:bl:remove")],
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:blacklist")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


@dp.callback_query(F.data == "admin:bl:add_username")
async def admin_bl_add_username(call: CallbackQuery, state: FSMContext):
    await state.update_data(bl_type="username")
    await call.message.edit_text(
        "🚫 <b>Добавление в ЧС</b>\n\n"
        "Шаг 1/2. Отправьте <b>@username</b>:\n\n"
        "Для отмены — /admin",
        parse_mode="HTML",
    )
    await state.set_state(Admin.blacklist_add_username)


@dp.message(Admin.blacklist_add_username)
async def admin_bl_save_username(message: Message, state: FSMContext):
    if message.text == "/admin":
        await state.clear()
        await show_admin_menu(message, state)
        return

    value = message.text.strip()
    if not value.startswith("@"):
        value = "@" + value

    await state.update_data(bl_value=value)
    await state.set_state(Admin.blacklist_reason)
    await message.answer("Шаг 2/2. Причина (или <code>-</code>):", parse_mode="HTML")


@dp.callback_query(F.data == "admin:bl:add_phone")
async def admin_bl_add_phone(call: CallbackQuery, state: FSMContext):
    await state.update_data(bl_type="phone")
    await call.message.edit_text(
        "🚫 <b>Добавление в ЧС</b>\n\n"
        "Шаг 1/2. Отправьте <b>номер телефона</b>:\n\n"
        "Для отмены — /admin",
        parse_mode="HTML",
    )
    await state.set_state(Admin.blacklist_add_phone)


@dp.message(Admin.blacklist_add_phone)
async def admin_bl_save_phone(message: Message, state: FSMContext):
    if message.text == "/admin":
        await state.clear()
        await show_admin_menu(message, state)
        return

    value = message.text.strip()
    await state.update_data(bl_value=value)
    await state.set_state(Admin.blacklist_reason)
    await message.answer("Шаг 2/2. Причина (или <code>-</code>):", parse_mode="HTML")


@dp.message(Admin.blacklist_reason)
async def admin_bl_save_reason(message: Message, state: FSMContext):
    reason = message.text.strip()
    if reason == "-":
        reason = ""

    data = await state.get_data()
    bl_type = data.get("bl_type", "username")
    bl_value = data.get("bl_value", "")

    ok = add_to_blacklist(bl_type, bl_value, reason)
    await state.clear()

    if ok:
        await message.answer(f"✅ Добавлено в ЧС: <code>{bl_value}</code>", parse_mode="HTML")
    else:
        await message.answer("⚠️ Уже в ЧС.")

    await show_admin_menu(message, state)


@dp.callback_query(F.data == "admin:bl:remove")
async def admin_bl_remove_start(call: CallbackQuery):
    if not blacklist:
        await call.answer("ЧС пуст", show_alert=True)
        return

    buttons = []
    for key, item in list(blacklist.items())[:30]:
        icon = "📱" if item.get("type") == "phone" else "👤"
        buttons.append([InlineKeyboardButton(
            text=f"🗑 {icon} {item.get('value')}",
            callback_data=f"admin:bl:del:{item.get('value')}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:blacklist")])

    await call.message.edit_text(
        "🗑 Выберите запись для удаления:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
    )


@dp.callback_query(F.data.startswith("admin:bl:del:"))
async def admin_bl_remove_confirm(call: CallbackQuery):
    value = call.data.split(":", 3)[3]
    ok = remove_from_blacklist(value)
    if ok:
        await call.answer("Удалено")
    else:
        await call.answer("Не найдено", show_alert=True)
    await admin_blacklist_menu(call)


# ================== ПОПЫТКИ РЕГИСТРАЦИИ (АДМИН) ==================
@dp.callback_query(F.data == "admin:reg_attempts")
async def admin_reg_attempts(call: CallbackQuery):
    if call.from_user.id != ADMIN_ID:
        await call.answer("Только для админа", show_alert=True)
        return

    text = get_reg_attempts_text(limit=30)[:4000]

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Обновить", callback_data="admin:reg_attempts")],
        [InlineKeyboardButton(text="⬅️ В админ-панель", callback_data="admin:back")],
    ])
    await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")


# ================== МЕРОПРИЯТИЯ (КЛИЕНТ) ==================
@dp.callback_query(F.data == "show_events")
async def show_events(call: CallbackQuery):
    now = datetime.now()
    upcoming = []
    for e in events:
        if not e.get("active", True):
            continue
        try:
            start = datetime.fromisoformat(e.get("start_dt") or "")
        except (ValueError, TypeError):
            print(f"[SHOW_EVENTS] Пропущено {e.get('id')} — битый start_dt={e.get('start_dt')!r}")
            continue
        if start > now:
            upcoming.append(e)

    if not upcoming:
        await call.message.answer("📭 Пока нет доступных событий.")
        await call.answer()
        return

    buttons = []
    for e in upcoming:
        buttons.append([InlineKeyboardButton(
            text=f"🎉 {e['title']} — {e['date_text']}",
            callback_data=f"event:{e['id']}"
        )])

    buttons.append([InlineKeyboardButton(
        text="🔄 Обновить", callback_data="show_events"
    )])
    buttons.append([InlineKeyboardButton(
        text="🏠 На главную", callback_data="main_menu"
    )])

    await call.message.answer(
        "🎉 <b>Выберите мероприятие:</b>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )
    await call.answer()



# ================== МОИ БИЛЕТЫ (КЛИЕНТ) ==================
@dp.callback_query(F.data == "my_tickets")
async def my_tickets(call: CallbackQuery):
    """Показывает список событий, где у пользователя есть билеты."""
    user_id = call.from_user.id

    all_user_tickets = [
        (code, t) for code, t in tickets.items()
        if str(t.get("user_id")) == str(user_id)
    ]

    if not all_user_tickets:
        await call.message.answer(
            "🎫 <b>Мои билеты</b>\n\n"
            "📭 У вас пока нет билетов.\n\n"
            "Нажмите «🎉 Мероприятия» чтобы купить.",
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
                [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
            ]),
        )
        return

    # Группируем по event_id
    by_event = {}
    for code, t in all_user_tickets:
        ev_id = str(t.get("event_id") or "")
        if not ev_id:
            continue
        if ev_id not in by_event:
            by_event[ev_id] = {
                "event_title": t.get("event_title", "—"),
                "event_date": t.get("event_date", "—"),
                "tickets": [],
            }
        by_event[ev_id]["tickets"].append((code, t))

    # Статистика
    active_cnt = sum(1 for _, t in all_user_tickets if not t.get("used") and not t.get("auto_used"))
    used_cnt = sum(1 for _, t in all_user_tickets if t.get("used") and not t.get("auto_used"))
    burned_cnt = sum(1 for _, t in all_user_tickets if t.get("auto_used"))

    text = (
        f"🎟 <b>Мои билеты</b>\n\n"
        f"Всего: <b>{len(all_user_tickets)}</b>\n"
        f"🟢 Активных: <b>{active_cnt}</b>\n"
        f"🔴 Использовано: <b>{used_cnt}</b>\n"
        f"⚪ Сгорело: <b>{burned_cnt}</b>\n\n"
        f"Выберите событие:"
    )

    buttons = []

    # Сортируем события по дате (сначала ближайшие)
    sorted_events = sorted(
        by_event.items(),
        key=lambda kv: kv[1]["event_date"] or "",
    )

    for ev_id, info in sorted_events:
        ev_tickets = info["tickets"]
        total = len(ev_tickets)
        ev_active = sum(1 for _, t in ev_tickets if not t.get("used") and not t.get("auto_used"))
        ev_used = sum(1 for _, t in ev_tickets if t.get("used") and not t.get("auto_used"))
        ev_burned = sum(1 for _, t in ev_tickets if t.get("auto_used"))

        # Иконка статуса
        if ev_active > 0:
            icon = "🟢"
        elif ev_used > 0:
            icon = "🔴"
        else:
            icon = "⚪"

        # Короткое название (до 30 символов)
        title_short = info["event_title"][:30]
        if len(info["event_title"]) > 30:
            title_short += "…"

        buttons.append([InlineKeyboardButton(
            text=f"{icon} {title_short} ({total})",
            callback_data=f"my_event:{ev_id}"
        )])

    buttons.append([InlineKeyboardButton(text="🔄 Обновить", callback_data="my_tickets")])

    # Архив билетов (если есть)
    archived_cnt = sum(1 for _, t in all_user_tickets if is_ticket_archived(t))
    if archived_cnt > 0:
        buttons.append([InlineKeyboardButton(
            text=f"📦 Архив билетов ({archived_cnt})",
            callback_data="my_tickets_archive"
        )])

    buttons.append([InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")])
    buttons.append([InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")])

    await call.message.answer(
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )

@dp.callback_query(F.data.startswith("my_event:"))
async def my_event_tickets(call: CallbackQuery):
    """Показывает билеты конкретного события."""
    event_id = call.data.split(":", 1)[1]
    user_id = call.from_user.id

    event = next((e for e in events if str(e["id"]) == str(event_id)), None)
    if not event:
        await call.answer("Событие не найдено", show_alert=True)
        return

    ev_tickets = [
        (code, t) for code, t in tickets.items()
        if str(t.get("user_id")) == str(user_id)
        and str(t.get("event_id")) == str(event_id)
    ]

    if not ev_tickets:
        await call.answer("Билетов на это событие нет", show_alert=True)
        return

    def sort_key(kv):
        _, t = kv
        if t.get("auto_used"):
            return (2, t.get("used_at") or "")
        if t.get("used"):
            return (1, t.get("used_at") or "")
        return (0, t.get("event_date") or "")

    ev_tickets.sort(key=sort_key)

    active = sum(1 for _, t in ev_tickets if not t.get("used") and not t.get("auto_used"))
    used = sum(1 for _, t in ev_tickets if t.get("used") and not t.get("auto_used"))
    burned = sum(1 for _, t in ev_tickets if t.get("auto_used"))

    text = (
        f"🎉 <b>{event['title']}</b>\n"
        f"📅 {event['date_text']}\n\n"
        f"🎫 Билетов: <b>{len(ev_tickets)}</b>\n"
        f"🟢 Активных: <b>{active}</b>\n"
        f"🔴 Погашено: <b>{used}</b>\n"
        f"⚪ Сгорело: <b>{burned}</b>\n\n"
    )

    buttons = []
    for code, t in ev_tickets:
        status = get_ticket_status_icon(t)
        type_label = "Парный" if t.get("type") == "pair" else "Обычный"
        issued_at = t.get("issued_at", "—")
        holder = t.get("holder", "—")

        paid_note = ""
        if t.get("type") == "pair" and t.get("partner_code"):
            partner_ticket = tickets.get(t.get("partner_code"))
            if partner_ticket:
                partner_uid = partner_ticket.get("user_id")
                partner_user = users.get(int(partner_uid or 0), {})
                partner_uname = partner_user.get("username") or ""
                if partner_uname and str(partner_ticket.get("user_id")) == str(user_id):
                    paid_note = f" (оплачен @{partner_uname})"

        text += (
            f"{status} <code>{code}</code> · {type_label}{paid_note}\n"
            f"   👤 {holder}\n"
            f"   🛒 {issued_at}\n\n"
        )

        if not t.get("used") and not t.get("auto_used"):
            suffix = get_button_suffix(t)
            buttons.append([InlineKeyboardButton(
                text=f"📱 Показать QR — {code[-8:]}{suffix}",
                callback_data=f"my_ticket:{code}"
            )])
            buttons.append([InlineKeyboardButton(
                text=f"📤 На вход — {code[-8:]}",
                callback_data=f"my_ticket:send:{code}"
            )])
            buttons.append([InlineKeyboardButton(
                text=f"🎁 Передать — {code[-8:]}",
                callback_data=f"my_ticket:transfer:{code}"
            )])

    buttons.append([InlineKeyboardButton(text="⬅️ К списку событий", callback_data="my_tickets")])
    buttons.append([InlineKeyboardButton(text="🔄 Обновить", callback_data=f"my_event:{event_id}")])
    buttons.append([InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")])

    await call.message.answer(
        text[:4000],
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )
    await call.answer()


# ================== АРХИВ БИЛЕТОВ (КЛИЕНТ) ==================
@dp.callback_query(F.data == "my_tickets_archive")
async def my_tickets_archive(call: CallbackQuery):
    """Показывает архив билетов гостя (старше 30 дней после авто-погашения)."""
    user_id = call.from_user.id

    archived = [
        (code, t) for code, t in tickets.items()
        if str(t.get("user_id")) == str(user_id) and is_ticket_archived(t)
    ]

    if not archived:
        await call.message.answer(
            "📦 <b>Архив билетов</b>\n\n"
            "📭 Архив пуст.\n\n"
            "<i>Сюда попадают сгоревшие билеты через 30 дней.</i>",
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ К билетам", callback_data="my_tickets")],
            ]),
        )
        return

    # ✅ Защита от None
    archived.sort(key=lambda kv: kv[1].get("used_at") or "", reverse=True)

    text = f"📦 <b>Архив билетов</b> ({len(archived)})\n\n"
    for code, t in archived[:30]:
        text += (
            f"⚪ <code>{code}</code>\n"
            f"   🎉 {t.get('event_title', '—')}\n"
            f"   📅 {t.get('event_date', '—')}\n"
            f"   🛒 Куплен: {t.get('issued_at', '—')}\n"
            f"   🕐 Сгорел: {t.get('used_at', '—')}\n\n"
        )

    if len(archived) > 30:
        text += f"... и ещё {len(archived) - 30}"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ К билетам", callback_data="my_tickets")],
        [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
    ])

    await call.message.answer(text[:4000], reply_markup=kb, parse_mode="HTML")




# ================== ОТПРАВКА БИЛЕТА НА ВХОД ==================
@dp.callback_query(F.data.startswith("my_ticket:send:"))
async def my_ticket_send_start(call: CallbackQuery):
    """Начало отправки билета checker'у."""
    code = call.data.split(":", 2)[2]
    user_id = call.from_user.id

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if str(ticket.get("user_id")) != str(user_id):
        await call.answer("Это не ваш билет", show_alert=True)
        return

    if ticket.get("auto_used"):
        await call.answer("Этот билет сгорел — вы не пришли", show_alert=True)
        return

    if ticket.get("used"):
        await call.answer("Этот билет уже использован", show_alert=True)
        return

    event_id = ticket.get("event_id")

    available = []
    for chk_id, c in checkers.items():
        if not c.get("active", True):
            continue
        if not c.get("user_id"):
            continue
        if not checker_can_access_event(c, event_id):
            continue
        available.append((chk_id, c))

    if not available:
        await call.answer()
        await call.message.answer(
            "⚠️ <b>Нет доступных проверяющих</b>\n\n"
            "Обратитесь к организатору.",
            parse_mode="HTML",
        )
        return

    if len(available) == 1:
        chk_id, c = available[0]
        await _do_send_ticket_to_checker(call, code, chk_id, c)
        return

    buttons = []
    for chk_id, c in available:
        name = c.get("name") or "проверяющий"
        buttons.append([InlineKeyboardButton(
            text=f"👤 {name}",
            callback_data=f"my_ticket:sent:{code}:{chk_id}"
        )])
    buttons.append([InlineKeyboardButton(
        text="⬅️ Назад", callback_data="my_tickets"
    )])

    await call.message.answer(
        f"📤 <b>Отправка билета на вход</b>\n\n"
        f"Код: <code>{code}</code>\n\n"
        f"Выберите, кому отправить:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="HTML",
    )
    await call.answer()

@dp.callback_query(F.data.startswith("my_ticket:transfer:"))
async def my_ticket_transfer_start(call: CallbackQuery, state: FSMContext):
    code = call.data.split(":", 2)[2]
    user_id = call.from_user.id

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return
    if str(ticket.get("user_id")) != str(user_id):
        await call.answer("Это не ваш билет", show_alert=True)
        return
    if ticket.get("auto_used"):
        await call.answer("Билет сгорел — передать нельзя", show_alert=True)
        return
    if ticket.get("used"):
        await call.answer("Билет уже использован — передать нельзя", show_alert=True)
        return

    event = next((e for e in events if str(e["id"]) == str(ticket.get("event_id"))), None)
    if not event:
        await call.answer("Событие не найдено", show_alert=True)
        return
    try:
        start_dt = datetime.fromisoformat(event.get("start_dt") or "")
    except (ValueError, TypeError):
        await call.answer("Ошибка в дате события", show_alert=True)
        return

    now = datetime.now()
    hours_left = (start_dt - now).total_seconds() / 3600

    if hours_left < 5:
        await call.answer(
            f"❌ Передача невозможна.\n\n"
            f"До начала: {hours_left:.1f} ч.\n"
            f"Передача разрешена не позднее чем за 5 часов.",
            show_alert=True,
        )
        return

    await state.update_data(transfer_code=code)
    await call.message.answer(
        f"🎁 <b>Передача билета</b>\n\n"
        f"Код: <code>{code}</code>\n"
        f"🎉 {event.get('title', '—')}\n"
        f"📅 {event.get('date_text', '—')}\n"
        f"⏱ До начала: <b>{hours_left:.1f} ч</b>\n\n"
        f"Отправьте <b>@username</b> получателя.\n"
        f"Он должен быть зарегистрирован в боте.\n\n"
        f"Для отмены — /start",
        parse_mode="HTML",
    )
    await state.set_state(Reg.transfer_input)
    await call.answer()


@dp.message(Reg.transfer_input)
async def my_ticket_transfer_input(message: Message, state: FSMContext):
    text = message.text.strip()
    if not text.startswith("@"):
        await message.answer("❌ Формат: <b>@username</b>", parse_mode="HTML")
        return

    data = await state.get_data()
    code = data.get("transfer_code")
    ticket = tickets.get(code)
    if not ticket:
        await message.answer("❌ Билет не найден.")
        await state.clear()
        return

    if str(ticket.get("user_id")) != str(message.from_user.id):
        await message.answer("❌ Это не ваш билет.")
        await state.clear()
        return

    event = next((e for e in events if str(e["id"]) == str(ticket.get("event_id"))), None)
    if not event:
        await message.answer("❌ Событие не найдено.")
        await state.clear()
        return
    try:
        start_dt = datetime.fromisoformat(event.get("start_dt") or "")
    except (ValueError, TypeError):
        await message.answer("❌ Ошибка в дате события.")
        await state.clear()
        return
    if (start_dt - datetime.now()).total_seconds() / 3600 < 5:
        await message.answer(
            "❌ Передача невозможна — до события осталось менее 5 часов.",
        )
        await state.clear()
        return

    receiver_uname = text.lower().lstrip("@")
    receiver_id = username_index.get(receiver_uname)

    if not receiver_id:
        await message.answer(
            f"❌ Пользователь <b>{text}</b> не найден в боте.\n"
            f"Он должен сначала зарегистрироваться.",
            parse_mode="HTML",
        )
        return

    if receiver_id == message.from_user.id:
        await message.answer("❌ Нельзя передать билет самому себе.")
        return

    receiver = users.get(receiver_id, {})
    if not receiver:
        await message.answer("❌ Получатель не зарегистрирован.")
        return

    sender = users.get(message.from_user.id, {})
    sender_name = f"{sender.get('name', '')} {sender.get('surname', '')}".strip() or f"@{message.from_user.username or '—'}"
    receiver_name = f"{receiver.get('name', '')} {receiver.get('surname', '')}".strip() or f"@{receiver_uname}"

    ticket["user_id"] = receiver_id
    ticket["holder"] = receiver_name
    ticket["issued_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    save_ticket(code)

    save_ticket_transfer({
        "ticket_code": code,
        "from_user_id": message.from_user.id,
        "from_username": message.from_user.username or "",
        "from_name": sender_name,
        "to_user_id": receiver_id,
        "to_username": receiver_uname,
        "to_name": receiver_name,
        "event_id": ticket.get("event_id", ""),
        "event_title": ticket.get("event_title", ""),
        "transferred_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "note": "",
    })

    await state.clear()

    try:
        await bot.send_message(
            receiver_id,
            f"🎁 <b>Вам передали билет!</b>\n\n"
            f"Код: <code>{code}</code>\n"
            f"🎉 {ticket.get('event_title', '—')}\n"
            f"📅 {ticket.get('event_date', '—')}\n"
            f"👤 От: {sender_name}\n\n"
            f"Откройте «🎟 Мои билеты», чтобы посмотреть QR.",
            parse_mode="HTML",
        )
    except Exception:
        pass

    await message.answer(
        f"✅ <b>Билет передан</b>\n\n"
        f"Код: <code>{code}</code>\n"
        f"Получатель: <b>{receiver_name}</b> (@{receiver_uname})\n\n"
        f"Билет больше не в вашем списке.",
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("my_ticket:sent:"))
async def my_ticket_send_to_checker(call: CallbackQuery):
    """Отправка билета конкретному checker'у."""
    parts = call.data.split(":")
    try:
        code = parts[2]
        chk_id = parts[3]
    except IndexError:
        await call.answer("Ошибка данных", show_alert=True)
        return

    user_id = call.from_user.id

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if str(ticket.get("user_id")) != str(user_id):
        await call.answer("Это не ваш билет", show_alert=True)
        return

    if ticket.get("auto_used"):
        await call.answer("Этот билет сгорел", show_alert=True)
        return

    if ticket.get("used"):
        await call.answer("Билет уже использован", show_alert=True)
        return

    checker = checkers.get(chk_id)
    if not checker:
        await call.answer("Проверяющий недоступен", show_alert=True)
        return

    if not checker_can_access_event(checker, ticket.get("event_id")):
        await call.answer("Этот проверяющий не работает с вашим событием", show_alert=True)
        return

    await _do_send_ticket_to_checker(call, code, chk_id, checker)


async def _do_send_ticket_to_checker(call: CallbackQuery, code: str, chk_id: str, checker: dict):
    """Общая функция: отправляет карточку билета checker'у."""
    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    owner_id = ticket.get("user_id")
    holder_name = ticket.get("holder", "—")

    holder_user = {}
    try:
        if owner_id:
            holder_user = users.get(int(owner_id), {})
    except (TypeError, ValueError):
        pass

    holder_username = holder_user.get("username") or ""
    holder_phone = holder_user.get("phone") or "—"

    text = (
        f"🎫 <b>Билет на проверку</b>\n\n"
        f"От: <b>{holder_name}</b>"
    )
    if holder_username:
        text += f" (@{holder_username})"
    text += (
        f"\nКод: <code>{code}</code>\n\n"
        f"👤 {holder_name}\n"
    )
    if holder_username:
        text += f"💬 @{holder_username}\n"
    text += f"📱 {holder_phone}\n\n"
    text += (
        f"🎟 {_ticket_type_label(ticket)}\n"
        f"🎉 {ticket.get('event_title', '—')}\n"
        f"📅 {ticket.get('event_date', '—')}\n"
        f"📌 Выдан: {ticket.get('issued_at', '—')}\n"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text="✅ Погасить билет",
            callback_data=f"checker:use:{code}"
        )],
        [InlineKeyboardButton(
            text="❌ Отклонить",
            callback_data=f"checker:reject:{code}"
        )],
    ])

    checker_user_id = checker.get("user_id")

    try:
        await bot.send_message(
            checker_user_id,
            text,
            reply_markup=kb,
            parse_mode="HTML",
        )

        checker_name = checker.get("name") or "проверяющий"
        await call.message.answer(
            f"✅ <b>Билет отправлен на вход!</b>\n\n"
            f"Код: <code>{code}</code>\n"
            f"Проверяющий: <b>{checker_name}</b>\n\n"
            f"Покажите свой экран, если проверяющий попросит.",
            parse_mode="HTML",
        )
        await call.answer("Отправлено ✅")

    except Exception as e:
        print(f"[SEND] Ошибка отправки checker'у {chk_id}: {e}")
        await call.answer(
            "Не удалось отправить. Проверяющий не начинал чат с ботом?",
            show_alert=True,
        )


# ================== ПОКАЗ QR (ОБЩИЙ) ==================
@dp.callback_query(F.data.startswith("my_ticket:"))
async def show_my_ticket_qr(call: CallbackQuery):
    """Показывает QR-код конкретного билета."""
    code = call.data.split(":", 1)[1]
    user_id = call.from_user.id

    ticket = tickets.get(code)
    if not ticket:
        await call.answer("Билет не найден", show_alert=True)
        return

    if str(ticket.get("user_id")) != str(user_id):
        await call.answer("Это не ваш билет", show_alert=True)
        return

    if ticket.get("auto_used"):
        await call.answer("Этот билет сгорел — вы не пришли", show_alert=True)
        return

    if ticket.get("used"):
        await call.answer("Этот билет уже использован", show_alert=True)
        return

    event = next(
        (e for e in events if str(e["id"]) == str(ticket.get("event_id"))),
        None
    )

    qr_file = make_qr_image(code)
    type_label = "Парный" if ticket.get("type") == "pair" else "Обычный"

    caption = (
        f"🎫 <b>Ваш билет</b>\n\n"
        f"👤 {ticket.get('holder', '—')}\n"
        f"🎟 {type_label}\n"
        f"🎉 {ticket.get('event_title', '—')}\n"
        f"📅 {ticket.get('event_date', '—')}\n"
        f"🛒 Куплен: {ticket.get('issued_at', '—')}\n"
    )
    if event and event.get("place"):
        caption += f"📍 {event.get('place')}\n"
    caption += (
        f"\nКод: <code>{code}</code>\n\n"
        "Покажите QR на входе."
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")
    ]])

    await call.message.answer_photo(
        qr_file,
        caption=caption,
        reply_markup=kb,
        parse_mode="HTML",
    )
    await call.answer()


# ================== ВЫБОР СОБЫТИЯ КЛИЕНТОМ ==================
@dp.callback_query(F.data.startswith("event:"))
async def event_selected(call: CallbackQuery, state: FSMContext):
    eid = call.data.split(":")[1]
    event = next((e for e in events if e["id"] == eid), None)
    if not event:
        await call.answer("Событие не найдено")
        return

    is_admin = (call.from_user.id == ADMIN_ID)

    if not is_admin and call.from_user.id not in users:
        await call.answer("Сначала пройдите регистрацию 👤", show_alert=True)
        return

    if sales_closed(event):
        await call.answer("🛑 Продажа билетов закрыта", show_alert=True)
        return

    await state.update_data(event_id=eid, event=event)

    if is_admin:
        buttons = [[InlineKeyboardButton(
            text="⬅️ Закрыть",
            callback_data="admin:back"
        )]]
    else:
        buttons = [[InlineKeyboardButton(
            text=f"🎟 Обычный — {event['price_single']} ₽",
            callback_data="ticket:single"
        )]]
        if event.get("price_pair", 0) > 0:
            buttons.append([InlineKeyboardButton(
                text=f"🎟 Парный — {event['price_pair']} ₽",
                callback_data="ticket:pair"
            )])
        buttons.append([InlineKeyboardButton(
            text="🏠 На главную", callback_data="main_menu"
        )])

    caption = (
        f"🎉 <b>{event['title']}</b>\n\n"
        f"📅 {event['date_text']}\n"
        f"📍 {event['place']}\n\n"
        f"{event['description']}\n\n"
    )
    if is_admin:
        caption += "👁 <i>Режим просмотра (админ)</i>"
    else:
        caption += "Выберите тип билета:"

    kb = InlineKeyboardMarkup(inline_keyboard=buttons)

    if event.get("photo_id"):
        await call.message.answer_photo(
            event["photo_id"], caption=caption,
            reply_markup=kb, parse_mode="HTML"
        )
    else:
        await call.message.answer(caption, reply_markup=kb, parse_mode="HTML")

    await state.set_state(Reg.ticket_type)


# ========== КОНЕЦ ЧАСТИ 14 ==========

# ================== ОПЛАТА ЧЕРЕЗ ЮKASSA ==================
@dp.callback_query(F.data == "ticket:single", Reg.ticket_type)
async def ticket_single(call: CallbackQuery, state: FSMContext):
    try:
        data = await state.get_data()
        event = data.get("event")
        print(f"[TICKET_SINGLE] user={call.from_user.id}, event={event.get('id') if event else None}")

        if not event:
            await call.answer("Ошибка: событие не найдено", show_alert=True)
            return

        await state.update_data(ticket_type="single", price=event["price_single"])
        await call.answer()
        await go_to_payment(call.message, state)
    except Exception as e:
        print(f"[TICKET_SINGLE] Ошибка: {e}")
        await call.answer(f"Ошибка: {e}", show_alert=True)


@dp.callback_query(F.data == "ticket:pair", Reg.ticket_type)
async def ticket_pair(call: CallbackQuery, state: FSMContext):
    try:
        data = await state.get_data()
        event = data.get("event")
        print(f"[TICKET_PAIR] user={call.from_user.id}, event={event.get('id') if event else None}")

        if not event or event.get("price_pair", 0) <= 0:
            await call.answer("Парные недоступны", show_alert=True)
            return

        await state.update_data(ticket_type="pair", price=event["price_pair"])
        await call.message.answer(
            "🎟 Введите <b>@username</b> второго участника:",
            parse_mode="HTML",
        )
        await state.set_state(Reg.partner_input)
        await call.answer()
    except Exception as e:
        print(f"[TICKET_PAIR] Ошибка: {e}")
        await call.answer(f"Ошибка: {e}", show_alert=True)


@dp.message(Reg.partner_input)
async def partner_input(message: Message, state: FSMContext):
    text = message.text.strip()
    if not text.startswith("@") or len(text) < 3:
        await message.answer("❌ Формат: <b>@username</b>", parse_mode="HTML")
        return

    partner_id = username_index.get(text.lower().lstrip("@").strip())
    if partner_id == message.from_user.id:
        await message.answer("❌ Нельзя себя.")
        return

    if partner_id and partner_id in users:
        await state.update_data(partner_id=partner_id)
        partner = users[partner_id]
        await message.answer(
            f"✅ Найден: {partner['name']} {partner['surname']}",
        )
        await go_to_payment(message, state)
    else:
        invite_link = f"https://t.me/{BOT_USERNAME}?start=invite_{message.from_user.id}"
        await message.answer(
            f"❌ Не найден. Отправьте ссылку:\n<code>{invite_link}</code>\n\n"
            "После регистрации напишите /start заново.",
            parse_mode="HTML",
        )


async def go_to_payment(message: Message, state: FSMContext):
    """Создаёт платёж в ЮKassa."""
    data = await state.get_data()
    price = data.get("price", 0)
    ticket_type = data.get("ticket_type", "single")
    event = data.get("event", {})
    user_id = message.chat.id

    ticket_label = "Парный" if ticket_type == "pair" else "Обычный"
    description = f"{event.get('title', 'Билет')} — {ticket_label} билет"

    payment = create_yookassa_payment(
        amount=price,
        description=description,
        user_id=user_id,
        metadata={
            "event_id": data.get("event_id", ""),
            "ticket_type": ticket_type,
            "price": str(price),
        },
    )

    if not payment:
        await message.answer(
            "⚠️ Ошибка при создании платежа. Попробуйте позже или напишите в поддержку."
        )
        return

    payment_id = payment["payment_id"]
    confirmation_url = payment["confirmation_url"]

    pending_payments[payment_id] = {
        "user_id": user_id,
        "event_id": data.get("event_id", ""),
        "event_title": event.get("title", ""),
        "ticket_type": ticket_type,
        "price": price,
        "partner_id": data.get("partner_id"),
        "name": f"{data.get('name', '')} {data.get('surname', '')}".strip(),
        "phone": data.get("phone", ""),
        "email": data.get("email", ""),
    }

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"💳 Оплатить {price} ₽", url=confirmation_url)],
        [InlineKeyboardButton(text="🆘 Техподдержка", url=f"https://t.me/{settings.get('support', '@your_support').lstrip('@')}")],
        [InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu")],
        [InlineKeyboardButton(text="❌ Отменить", callback_data="cancel_order")],
    ])

    await message.answer(
        f"💳 <b>Билет: {ticket_label} — {price} ₽</b>\n\n"
        f"🎉 {event.get('title', '')}\n"
        f"📅 {event.get('date_text', '')}\n\n"
        "1. Нажмите «Оплатить» — откроется страница ЮKassa.\n"
        "2. Оплатите картой, СБП или другим способом.\n"
        "3. После оплаты билет придёт автоматически в течение 1 минуты.\n\n"
        "⏱ Страница оплаты активна 60 минут.\n\n"
        "<b>⚠️ Во время перехода на страницу Юкассы желательно отключить VPN, иначе ссылка для оплаты может не открыться.</b>",
        reply_markup=kb, parse_mode="HTML",
    )
    await state.set_state(Reg.payment)


@dp.callback_query(F.data == "cancel_order")
async def cancel_order(call: CallbackQuery, state: FSMContext):
    await state.clear()
    await call.message.edit_text("❌ Отменено. /start — заново.")




# ================== ВЫДАЧА БИЛЕТА ==================
async def issue_ticket(user_id: int, code: str, ticket_type: str, holder_name: str, event: dict):
    tickets[code] = {
        "user_id": user_id,
        "type": ticket_type,
        "holder": holder_name,
        "used": False,
        "used_at": "",
        "used_by": "",
        "auto_used": False,
        "archived_at": None,
        "event_id": event.get("id", ""),
        "event_title": event.get("title", ""),
        "event_date": event.get("date_text", ""),
        "issued_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "partner_code": "",
    }
    save_ticket(code)

    qr_file = make_qr_image(code)
    type_label = "Парный" if ticket_type == "pair" else "Обычный"

    holder_user = users.get(user_id, {})
    holder_username = holder_user.get("username") or ""

    caption = (
        f"🎫 <b>Ваш билет на {event.get('title', 'Vida Bachata')}</b>\n\n"
        f"👤 {holder_name}\n"
    )
    if holder_username:
        caption += f"💬 @{holder_username}\n"

    caption += (
        f"🎟 {type_label}\n"
        f"📅 {event.get('date_text', '')}\n"
        f"📍 {event.get('place', '')}\n"
        f"🛒 Куплен: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n\n"
        f"Код: <code>{code}</code>\n\n"
        "Покажите QR на входе."
    )

    kb_ticket = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets"),
        InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu"),
    ]])

    try:
        await bot.send_photo(
            user_id,
            qr_file,
            caption=caption,
            reply_markup=kb_ticket,
            parse_mode="HTML",
        )
    except Exception as e:
        print(f"[TICKET] Ошибка отправки билета user={user_id}: {e}")


async def process_successful_payment(payment_id: str):
    """Обрабатывает успешную оплату с защитой от двойной обработки."""
    if payment_id in processed_payments:
        print(f"[YOOKASSA] ⚠️ Платёж {payment_id} уже был обработан — пропускаем")
        return

    pending = pending_payments.pop(payment_id, None)
    if not pending:
        print(f"[YOOKASSA] payment_id={payment_id} не найден в pending")
        processed_payments.add(payment_id)
        return

    processed_payments.add(payment_id)
    print(f"[YOOKASSA] ✅ Платёж {payment_id} принят в обработку")

    user_id = pending["user_id"]
    event = next((e for e in events if e["id"] == pending["event_id"]), None)
    if not event:
        print(f"[YOOKASSA] Событие {pending['event_id']} не найдено")
        return

    ticket_type = pending["ticket_type"]
    buyer_user = users.get(user_id, {})
    buyer_full_name = " ".join(filter(None, [
        (buyer_user.get("name") or "").strip(),
        (buyer_user.get("surname") or "").strip(),
])).strip()

    buyer_name = (
        buyer_full_name
        or pending["name"]
        or (f"@{buyer_user.get('username', '')}" if buyer_user.get("username") else "Клиент")
)

    order = {
        "order_id": next_order_id(),
        "user_id": user_id,
        "event_id": pending["event_id"],
        "event_title": event.get("title", ""),
        "email": pending.get("email", ""),
        "name": buyer_name,
        "ticket_type": ticket_type,
        "price": pending["price"],
        "paid_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    orders.append(order)
    save_order(order)

    buyer_code = generate_ticket_code()
    await issue_ticket(user_id, buyer_code, ticket_type, buyer_name, event)

    if ticket_type == "pair" and pending.get("partner_id"):
        partner_id = pending["partner_id"]
        if partner_id in users:
            partner = users[partner_id]
            partner_name = f"{partner['name']} {partner['surname']}"
            partner_code = generate_ticket_code()
            await issue_ticket(
                partner_id, partner_code,
                ticket_type, partner_name, event
            )
            # Связываем билеты между собой
            tickets[buyer_code]["partner_code"] = partner_code
            tickets[partner_code]["partner_code"] = buyer_code
            save_ticket(buyer_code)
            save_ticket(partner_code)

    buyer_user = users.get(user_id, {})
    buyer_full_name = " ".join(filter(None, [
        (buyer_user.get("name") or "").strip(),
        (buyer_user.get("surname") or "").strip(),
    ])).strip()
    if not buyer_full_name:
        buyer_full_name = buyer_name

    buyer_username = buyer_user.get("username") or ""
    buyer_phone = buyer_user.get("phone") or pending.get("phone", "") or "—"
    ticket_type_label = "Парный" if ticket_type == "pair" else "Обычный"

       # Чек для самозанятого (ФНС)
    receipt_link = await send_fns_receipt(
        amount=pending["price"],
        description=f"Билет на {event.get('title', 'мероприятие')}",
        user_email=pending.get("email", ""),
    )

    admin_lines = [
        f"💰 <b>Оплата получена!</b>\n\n",
        f"🆔 Платёж: <code>{payment_id}</code>\n",
        f"🎉 {event.get('title', '')}\n",
        f"📅 {event.get('date_text', '')}\n",
        "\n",
        f"👤 <b>{buyer_full_name}</b>\n",
        f"🆔 <code>{user_id}</code>\n",
    ]
    if buyer_username:
        admin_lines.append(f"💬 @{buyer_username}\n")
    admin_lines.append(f"📱 {buyer_phone}\n")
    admin_lines.append("\n")
    admin_lines.append(f"🎟 {ticket_type_label} | {pending['price']} ₽\n")

    if receipt_link:
        admin_lines.append(f"\n🧾 Чек: {receipt_link}\n")

    try:
        await bot.send_message(ADMIN_ID, "".join(admin_lines), parse_mode="HTML")
    except Exception:
        pass

    client_lines = [
        "✅ <b>Оплата прошла!</b>\n\n",
        f"👤 {buyer_full_name}\n",
        f"🎉 {event.get('title', '')}\n",
        f"📅 {event.get('date_text', '')}\n",
        f"🛒 Куплен: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n",
    ]
    if event.get("place"):
        client_lines.append(f"📍 {event.get('place')}\n")
    client_lines.append(f"🎟 {ticket_type_label} | {pending['price']} ₽\n\n")
    client_lines.append("🎫 Билет отправлен выше ⬆️\n")
    client_lines.append("Покажите QR-код на входе.")
    if receipt_link:
        client_lines.append(f"\n🧾 Чек: {receipt_link}\n")

    kb_client = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets"),
        InlineKeyboardButton(text="🏠 На главную", callback_data="main_menu"),
    ]])

    try:
        await bot.send_message(
            user_id,
            "".join(client_lines),
            reply_markup=kb_client,
            parse_mode="HTML",
        )
    except Exception:
        pass

    if len(processed_payments) > 5000:
        to_remove = list(processed_payments)[:1000]
        for pid in to_remove:
            processed_payments.discard(pid)


# ========== КОНЕЦ ЧАСТИ 15 ==========


# ================== ГЛАВНОЕ МЕНЮ ==================
async def show_main_menu(call: CallbackQuery, state: FSMContext):
    """Возвращает пользователя в главное меню."""
    await state.clear()
    user_id = call.from_user.id

    if user_id == ADMIN_ID:
        text = (
            f"👋 Привет!\n"
            f"Вы вошли как <b>главный администратор</b>.\n\n"
            f"Используйте панель ниже для управления ботом.\n\n"
            f"Подпишитесь на соцсети:\n"
            f"ВК (https://vk.com/vidabachata) • "
            f"Telegram (https://t.me/+eTMoEG6V1883Nzli)"
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🛠 Админ-панель", callback_data="admin:back")],
            [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
            [InlineKeyboardButton(text="🎟 Мои билеты", callback_data="my_tickets")],
            [InlineKeyboardButton(text="ВК", url="https://vk.com/vidabachata"),
             InlineKeyboardButton(text="Telegram", url="https://t.me/+eTMoEG6V1883Nzli")],
        ])
    else:
        checker = find_checker_for_user(
            user_id=user_id,
            username=call.from_user.username or "",
        )
        if checker:
            return await checker_start(call.message, state, checker)
        else:
            name = ""
            if user_id in users:
                name = users[user_id].get("name", "")
            greeting = f"👋 Привет, <b>{name}</b>!\n\n" if name else "👋 Привет!\n\n"
            text = (
                f"{greeting}"
                f"Здесь можно купить билет на вечеринки <b>Vida Bachata</b> 💃🕺\n\n"
                f"Выберите действие:"
            )
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🎉 Мероприятия", callback_data="show_events")],
                [InlineKeyboardButton(text="🎫 Мои билеты", callback_data="my_tickets")],
                [InlineKeyboardButton(text="ВК", url="https://vk.com/vidabachata"),
                 InlineKeyboardButton(text="Telegram", url="https://t.me/+eTMoEG6V1883Nzli")],
            ])

    try:
        await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")
    except Exception:
        await call.message.answer(text, reply_markup=kb, parse_mode="HTML")
    await call.answer()


@dp.callback_query(F.data == "main_menu")
async def main_menu_callback(call: CallbackQuery, state: FSMContext):
    await show_main_menu(call, state)




# ================== WEBHOOK ДЛЯ RENDER ==================
_main_loop = None


def start_keepalive_server():
    """HTTP-сервер на порту + webhook ЮKassa."""
    PORT = int(os.getenv("PORT", "10000"))

    class MainHandler(http.server.BaseHTTPRequestHandler):
        def do_HEAD(self):
            self.send_response(200)
            self.send_header("Content-type", "text/plain")
            self.end_headers()


        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-type", "text/plain")
            self.end_headers()
            self.wfile.write(b"OK")

        def do_POST(self):
            if self.path == "/yookassa/webhook":
                try:
                    content_length = int(self.headers.get("Content-Length", 0))
                    body = self.rfile.read(content_length).decode("utf-8")
                    data = json.loads(body)

                    event = data.get("event", "")
                    obj = data.get("object", {})
                    payment_id = obj.get("id", "")
                    status = obj.get("status", "")

                    client_ip = self.client_address[0]
                    forwarded = self.headers.get("X-Forwarded-For", "")
                    print(
                        f"[YOOKASSA WEBHOOK] event={event}, payment_id={payment_id}, "
                        f"status={status}, ip={client_ip}, xff={forwarded}"
                    )

                    if event == "payment.succeeded" and status == "succeeded":
                        if _main_loop:
                            asyncio.run_coroutine_threadsafe(
                                verify_and_process(payment_id),
                                _main_loop
                            )
                        else:
                            print(
                                "[YOOKASSA WEBHOOK] ⚠️ _main_loop не установлен — "
                                "платёж проигнорирован"
                            )

                    self.send_response(200)
                    self.send_header("Content-type", "text/plain")
                    self.end_headers()
                    self.wfile.write(b"OK")

                except Exception as e:
                    print(f"[YOOKASSA WEBHOOK] Ошибка: {e}")
                    self.send_response(500)
                    self.end_headers()
            else:
                self.send_response(404)
                self.end_headers()

        def log_message(self, format, *args):
            pass

    class ReusableTCPServer(socketserver.ThreadingTCPServer):
        allow_reuse_address = True

    with ReusableTCPServer(("", PORT), MainHandler) as httpd:
        print(f"[KEEPALIVE] HTTP-сервер запущен на порту {PORT}")
        httpd.serve_forever()


async def verify_and_process(payment_id: str):
    """Проверяет платёж через API ЮKassa и обрабатывает если реальный."""
    if payment_id in processed_payments:
        print(f"[YOOKASSA] ⚠️ Платёж {payment_id} уже обработан (до верификации)")
        return

    pending = pending_payments.get(payment_id)
    if not pending:
        print(f"[YOOKASSA] ❌ Платёж {payment_id} не найден в pending — игнорируем")
        return

    is_real = await asyncio.to_thread(verify_payment_with_yookassa, payment_id)
    if not is_real:
        print(f"[YOOKASSA] ❌ Поддельный или неизвестный платёж {payment_id} — игнорируем")
        return

    try:
        payment = await asyncio.to_thread(Payment.find_one, payment_id)
        meta_user_id = None
        if payment and payment.metadata:
            meta_user_id = payment.metadata.get("user_id")
        if str(meta_user_id) != str(pending["user_id"]):
            print(
                f"[YOOKASSA] ❌ metadata.user_id={meta_user_id} "
                f"≠ pending.user_id={pending['user_id']} — игнорируем"
            )
            return
    except Exception as e:
        print(f"[YOOKASSA] ⚠️ Не удалось сверить metadata: {e}")
        return

    print(f"[YOOKASSA] ✅ Платёж {payment_id} подтверждён через API — обрабатываем")
    await process_successful_payment(payment_id)


async def _self_ping_loop():
    """Пингует свой URL каждые 10 минут."""
    render_url = os.getenv("RENDER_EXTERNAL_URL", "")
    if not render_url:
        print("[KEEPALIVE] RENDER_EXTERNAL_URL не задан — самопинг отключён")
        return

    print(f"[KEEPALIVE] Самопинг активирован для {render_url}")
    while True:
        await asyncio.sleep(600)
        try:
            await asyncio.to_thread(urllib.request.urlopen, render_url, timeout=10)
            print("[KEEPALIVE] Самопинг успешен")
        except Exception as e:
            print(f"[KEEPALIVE] Самопинг ошибка: {e}")


def start_keepalive():
    """Запускает HTTP-сервер в фоновом потоке."""
    server_thread = threading.Thread(target=start_keepalive_server, daemon=True)
    server_thread.start()


def start_self_ping():
    """Запускает самопинг в фоне (async)."""
    asyncio.create_task(_self_ping_loop())


# ================== ЗАПУСК ==================
async def main():
    global _main_loop
    _main_loop = asyncio.get_running_loop()

    load_data()
    load_transfers()
    print("🤖 Бот запускается...")

    start_keepalive()
    start_self_ping()
    start_auto_burn()

    me = await bot.get_me()
    print(f"✅ Бот @{me.username} работает")

    # Сбрасываем старые getUpdates и webhook перед polling
    try:
        await bot.delete_webhook(drop_pending_updates=True)
        print("✅ Webhook удалён, старые обновления сброшены")
    except Exception as e:
        print(f"⚠️ delete_webhook: {e}")

    # Пауза, чтобы Telegram успел «забыть» старый инстанс
    await asyncio.sleep(2)

    print("🚀 Запускаю polling...")
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
